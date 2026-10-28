"""P7/P8：内部XLSX闭环，严格结构校验；永不生成百度执行表头。"""
from __future__ import annotations

import io
import json
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.sem_phase1a.contracts import PlanRow, WARNING

MAX_XLSX_BYTES = 2 * 1024 * 1024
HEADERS = ["内部行ID", "内部父行ID", "内部类型", "内部内容（可改）", "内部预算分（可改）",
           "核实状态", "source_refs", "顾问演练复核（只读）"]
SHEETS = ["内部规划", "来源", "演练绑定"]


def _meta(view, audience):
    return [["schema", "sem-shadow-slice1-v1"], ["warning", WARNING], ["plan_id", view["plan_id"]],
            ["version", view["version"]], ["plan_sha256", view["plan_sha256"]],
            ["submission_sha256", view["submission_sha256"]], ["audience", audience],
            ["mode", "rehearsal"], ["execution_export_enabled", False], ["customer_confirmation", "pending"]]


def _source_rows(view):
    return [[s["source_id"], s["field"], s["value"], s["origin"], s["artifact_sha256"], s["locator"],
             "顾问已裁定" if s["source_id"] in view["selected"].values() else "候选来源"] for s in view["sources"]]


def _verification(row):
    return "待核实" if row["verification"] == "pending_verification" else "已裁定来源"


def _review_row(row, view):
    return [row["row_id"], row["parent_id"], row["kind"], row["text"], row["budget_cents"],
            _verification(row), json.dumps(row["source_refs"], ensure_ascii=False),
            "已演练复核" if row["row_id"] in view["approvals"] else "待顾问逐项复核"]


def _literal(cell, value):
    cell.value = value
    # 不让任何来源/编辑文字成为公式或外部DDE引用。
    if isinstance(value, str):
        cell.data_type = "s"


def _append(sheet, values):
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        _literal(sheet.cell(row, column), value)


def write_workbook(view, audience):
    if audience not in {"advisor", "customer"}:
        raise ValueError("仅顾问或客户演练样例")
    book = Workbook()
    sheet = book.active
    sheet.title = SHEETS[0]
    _append(sheet, [WARNING])
    _append(sheet, ["客户审核表样例（仅演练）" if audience == "customer" else "顾问内部规划"])
    _append(sheet, ["仅修改内容/预算列；导入先预览，顾问在服务端逐项复核。所有推断保持待核实。"])
    _append(sheet, HEADERS)
    for row in view["rows"]:
        _append(sheet, _review_row(row, view))
    source_sheet = book.create_sheet(SHEETS[1])
    _append(source_sheet, [WARNING])
    _append(source_sheet, ["来源ID", "字段", "合成原文", "来源类型", "原文SHA-256", "定位/版本", "裁定"])
    for row in _source_rows(view):
        _append(source_sheet, row)
    binding = book.create_sheet(SHEETS[2])
    for row in _meta(view, audience):
        _append(binding, row)
    widths = ([18, 18, 34, 70, 24, 22, 30, 28], [18, 24, 70, 24, 70, 100, 24], [30, 85])
    for ws, cols in zip(book.worksheets, widths):
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=11, color="172033")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for index, width in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = width
        ws.row_dimensions[1].height = 30
        ws.cell(1, 1).font = Font(name="Arial", size=14, bold=True, color="B42318")
        ws.cell(1, 1).alignment = Alignment(vertical="center", wrap_text=False)
        ws.oddHeader.center.text = WARNING
        ws.oddHeader.center.size = 10
    sheet.cell(3, 1).alignment = Alignment(vertical="center", wrap_text=False)
    sheet.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=False)
    for ws, header in ((sheet, 4), (source_sheet, 2)):
        for cell in ws[header]:
            cell.fill = PatternFill("solid", fgColor="243954")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[header].height = 32
        ws.freeze_panes = f"D{header+1}" if ws == sheet else f"C{header+1}"
        ws.auto_filter.ref = f"A{header}:{ws.cell(ws.max_row, len(widths[0] if ws == sheet else widths[1])).coordinate}"
        for row in range(header+1, ws.max_row+1):
            ws.row_dimensions[row].height = 58 if ws == source_sheet else 46
    sheet.row_dimensions[3].height = 34
    for row in range(5, sheet.max_row+1):
        for col in (4, 5):
            sheet.cell(row, col).fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.cell(row, 5).number_format = '#,##0"分"'
    sheet.print_title_rows = "1:4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _safe_archive(data):
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ValueError("内部Excel大小超限或为空")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(i.file_size for i in members) > 10 * 1024 * 1024:
                raise ValueError("内部Excel解压大小超限")
            for item in members:
                name = item.filename.lower()
                if (item.flag_bits & 1 or ".." in name or "vba" in name or "externallink" in name
                        or name.startswith(("/", "\\"))):
                    raise ValueError("内部Excel含不允许的成员")
                if name.endswith((".xml", ".rels")):
                    content = archive.read(item).lower()
                    if any(token in content for token in (b"<!doctype", b"<!entity", b'targetmode="external"',
                                                           b"targetmode='external'")):
                        raise ValueError("内部Excel禁止外部引用/实体")
    except (zipfile.BadZipFile, RuntimeError, OSError):
        raise ValueError("不是有效内部XLSX") from None


def parse_workbook(data, view, audience):
    _safe_archive(data)
    try:
        book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception:
        raise ValueError("内部Excel无法读取") from None
    try:
        if audience not in {"advisor", "customer"} or book.sheetnames != SHEETS:
            raise ValueError("内部Excel页签或用途不匹配")
        for ws in book:
            if ws.max_row > 1100 or ws.max_column > 8:
                raise ValueError("内部Excel维度超限")
            for row in ws:
                if any(cell.data_type == "f" or cell.hyperlink for cell in row):
                    raise ValueError("内部Excel禁止公式/超链接")
        meta = [list(r) for r in book[SHEETS[2]].values]
        # openpyxl可能把尾部空列补齐；严格检查且不容许重复绑定字段。
        if meta != _meta(view, audience):
            raise ValueError("演练规划ID/版本/hash/用途绑定不匹配")
        sources = [list(r) for r in book[SHEETS[1]].values]
        expected_sources = [[WARNING] + [None]*6,
                            ["来源ID", "字段", "合成原文", "来源类型", "原文SHA-256", "定位/版本", "裁定"]] + _source_rows(view)
        if sources != expected_sources:
            raise ValueError("来源页不可篡改")
        ws = book[SHEETS[0]]
        if (ws.max_row != len(view["rows"])+4 or ws.max_column != len(HEADERS)
                or ws.cell(1, 1).value != WARNING or [c.value for c in ws[4]] != HEADERS):
            raise ValueError("内部表头/行数量不匹配，禁止增删行")
        result = []
        for index, old in enumerate(view["rows"], 5):
            values = [c.value for c in ws[index]]
            expected = _review_row(old, view)
            if any(values[col] != expected[col] for col in (0, 1, 2, 5, 6, 7)):
                raise ValueError("只允许修改内容和预算列，不可篡改身份、来源、核实或批准")
            if not isinstance(values[3], str) or not values[3].strip():
                raise ValueError("内部内容须为非空文本")
            updated = {**old, "text": values[3], "budget_cents": values[4]}
            try:
                row = PlanRow.model_validate(updated)
            except Exception:
                raise ValueError("内部行内容/整数分预算不满足契约") from None
            if (row.text, row.budget_cents) != (old["text"], old["budget_cents"]):
                updated = row.model_dump()
                updated["verification"] = "pending_verification"
                result.append(updated)
            else:
                result.append(row.model_dump())
        return result
    finally:
        book.close()
