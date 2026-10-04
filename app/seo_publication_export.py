"""Monthly SEO publication list, independent of HTTP and database access."""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image


BEIJING = timezone(timedelta(hours=8))
DEFAULT_PUBLICATION_LIST_TEMPLATE = (
    {"key": "number", "title": "序号", "width": 8, "type": "text"},
    {"key": "platform", "title": "平台/渠道", "width": 18, "type": "text"},
    {"key": "title", "title": "文章标题", "width": 42, "type": "text"},
    {"key": "keywords", "title": "关键词", "width": 30, "type": "text"},
    {"key": "page_url", "title": "发布链接", "width": 48, "type": "link"},
    {"key": "published_at", "title": "发布时间（北京时间）", "width": 23, "type": "datetime"},
    {"key": "capture_status", "title": "截图状态", "width": 32, "type": "text"},
    {"key": "image_key", "title": "截图缩略图", "width": 31, "type": "image"},
    {"key": "captured_at", "title": "截图时间（北京时间）", "width": 23, "type": "datetime"},
    {"key": "notes", "title": "备注", "width": 40, "type": "text"},
)

CAPTURE_ERRORS = {
    "capture_disabled": "截图功能未开启", "capture_recent": "刚提交过请稍后",
    "invalid_site_url": "链接不属于该站点", "timeout": "截图超时",
    "publication_url_missing": "该发布记录没有已登记的发布链接",
    "publication_url_mismatch": "链接与发布记录不一致",
    "captcha_page": "目标平台要求人机验证，未能截取正文",
    "blocked_by_platform": "目标平台拒绝了自动访问（可能需要登录）",
    "invalid_image": "图片文件无法识别", "unsupported_image_type": "仅支持 PNG、JPG、WebP 图片",
    "image_too_large": "图片过大",
}


def month_bounds(month: str) -> tuple[datetime, datetime]:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise ValueError("月份须为 YYYY-MM")
    year, number = map(int, month.split("-"))
    start = datetime(year, number, 1, tzinfo=BEIJING)
    end = datetime(year + (number == 12), number % 12 + 1, 1, tzinfo=BEIJING)
    return start.astimezone(timezone.utc).replace(tzinfo=None), end.astimezone(timezone.utc).replace(tzinfo=None)


def beijing_text(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.replace(tzinfo=timezone.utc).astimezone(BEIJING).strftime("%Y-%m-%d %H:%M") if value.tzinfo is None else value.astimezone(BEIJING).strftime("%Y-%m-%d %H:%M")


def capture_fields(captures: list) -> dict:
    """Latest successful evidence wins; a newer failure is retained in notes."""
    ordered = sorted(captures, key=lambda c: (c.captured_at, c.id), reverse=True)
    success = next((c for c in ordered if c.status == "succeeded"), None)
    latest = ordered[0] if ordered else None
    if success:
        status = "人工上传" if success.source == "manual" else "自动截图"
        recent_failure = next((c for c in ordered if c.status == "failed" and (c.captured_at, c.id) > (success.captured_at, success.id)), None)
        note = f"最近一次截图失败：{CAPTURE_ERRORS.get(recent_failure.error_code, recent_failure.error_code or '截图失败')}" if recent_failure else ""
        return {"capture_status": status, "image_key": success.storage_key, "captured_at": success.captured_at, "capture_note": note, "capture_kind": "manual" if success.source == "manual" else "auto"}
    if latest and latest.status == "failed":
        reason = CAPTURE_ERRORS.get(latest.error_code, latest.error_code or "截图失败")
        return {"capture_status": f"失败：{reason}", "image_key": None, "captured_at": None, "capture_note": "", "capture_kind": "failed"}
    return {"capture_status": "无截图", "image_key": None, "captured_at": None, "capture_note": "", "capture_kind": "none"}


def publication_summary(rows: list[dict], site_name: str, month: str, exported_at: datetime) -> dict:
    kinds = Counter(row["capture_kind"] for row in rows)
    return {"site": site_name, "month": month, "exported_at": exported_at, "total": len(rows),
            "platforms": Counter(row.get("platform") or "未记录" for row in rows),
            "auto": kinds["auto"], "manual": kinds["manual"], "failed": kinds["failed"], "none": kinds["none"]}


def _safe_text(value) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _thumbnail(data: bytes, max_pixels: int) -> bytes:
    Image.MAX_IMAGE_PIXELS = max_pixels
    with Image.open(BytesIO(data)) as source:
        if source.width * source.height > max_pixels:
            raise ValueError("图片过大")
        top_height = min(source.height, max(1, round(source.width * 1.2)))
        image = source.crop((0, 0, source.width, top_height)).convert("RGB")
        image.thumbnail((200, 240))
        output = BytesIO()
        image.save(output, format="JPEG", quality=75, optimize=True)
        return output.getvalue()


def build_publication_list_workbook(rows: list[dict], summary: dict, template=DEFAULT_PUBLICATION_LIST_TEMPLATE,
                                    image_loader=None, *, max_pixels: int = 16_000_000) -> bytes:
    counts = {key: summary[key] for key in ("auto", "manual", "failed", "none")}
    workbook = Workbook()
    detail = workbook.active
    detail.title = "发布明细"
    detail.append([column["title"] for column in template])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:{detail.cell(1, len(template)).column_letter}{len(rows) + 1}"
    for index, column in enumerate(template, 1):
        detail.column_dimensions[detail.cell(1, index).column_letter].width = column["width"]
        header = detail.cell(1, index)
        header.fill = PatternFill("solid", fgColor="1E4770")
        header.font = Font(color="FFFFFF", bold=True)
    for row_number, record in enumerate(rows, 2):
        detail.row_dimensions[row_number].height = 22
        image_error = False
        for column_number, spec in enumerate(template, 1):
            cell = detail.cell(row_number, column_number)
            key, kind = spec["key"], spec["type"]
            if kind == "image":
                if record.get(key) and image_loader:
                    try:
                        data = _thumbnail(image_loader(record[key]), max_pixels)
                        picture = SheetImage(BytesIO(data))
                        picture.anchor = cell.coordinate
                        detail.add_image(picture)
                        detail.row_dimensions[row_number].height = 185
                    except Exception:
                        image_error = True
                elif record.get("capture_kind") in {"auto", "manual"}:
                    image_error = True
                continue
            value = row_number - 1 if key == "number" else record.get(key)
            if kind == "datetime":
                value = beijing_text(value)
            cell.value = value if key == "number" else _safe_text(value)
            cell.alignment = Alignment(vertical="center", wrap_text=key in {"title", "notes", "capture_status"})
            if kind == "link" and isinstance(value, str) and value.startswith(("https://", "http://")):
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
        if image_error:
            status_cell = next((detail.cell(row_number, n) for n, col in enumerate(template, 1) if col["key"] == "capture_status"), None)
            if status_cell is not None:
                status_cell.value = "失败：截图图片不可用"
            kind = record.get("capture_kind")
            if kind in {"auto", "manual"}:
                counts[kind] -= 1
                counts["failed"] += 1
    info = workbook.create_sheet("汇总")
    covered = counts["auto"] + counts["manual"]
    items = [("站点", summary["site"]), ("月份", summary["month"]),
             ("导出时间（北京时间）", beijing_text(summary["exported_at"])),
             ("发布总数", summary["total"]), ("截图覆盖数", covered),
             ("截图覆盖率", covered / summary["total"] if summary["total"] else 0),
             ("自动截图数", counts["auto"]), ("人工上传数", counts["manual"]),
             ("截图失败数", counts["failed"]), ("无截图数", counts["none"])]
    items.extend((f"平台/渠道：{name}", count) for name, count in sorted(summary["platforms"].items()))
    for label, value in items:
        info.append((_safe_text(label), value))
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 42
    info["B6"].number_format = "0.0%"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
