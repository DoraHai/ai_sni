"""Build real XLSX and HTML from the same scoped GEO evidence snapshot."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from html import escape
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.geo.content.report_data import citation_rows, local_time, provenance, trend_rows

SECTION_DEFAULTS = [
    {"key": "overview", "title": "概览指标", "visible": True},
    {"key": "trends", "title": "各 AI 引擎提及率与趋势", "visible": True},
    {"key": "citations", "title": "发布 URL 引用", "visible": True},
    {"key": "publications", "title": "发布清单", "visible": True},
    {"key": "provenance", "title": "样本来源与口径", "visible": True},
    {"key": "appendix", "title": "问题与引擎明细", "visible": True},
]
SOURCE_NAMES = {"real": "真实引擎采样", "manual": "人工录入", "simulated": "模拟/演示", "unknown": "来源未知"}


def report_template(profile: dict | None) -> list[dict]:
    raw = (profile or {}).get("geo_report_template")
    if not isinstance(raw, list) or {x.get("key") for x in raw if isinstance(x, dict)} != {x["key"] for x in SECTION_DEFAULTS}:
        return [dict(x) for x in SECTION_DEFAULTS]
    return [{"key": x["key"], "title": str(x.get("title") or x["key"])[:60], "visible": bool(x.get("visible", True))} for x in raw]


def safe_cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value
    return value


def build_xlsx(data: dict, *, sample_label: str = "") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, headers):
        ws = wb.create_sheet(name)
        ws.append([safe_cell(sample_label if sample_label else "GEO 真实记录；空白表示无数据")])
        ws.append(headers)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="245277")
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{ws.cell(2, len(headers)).column_letter}2"
        return ws
    def format_columns(ws, widths, wrap_columns=()):
        for column, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(2, column).column_letter].width = width
        wrapped = Alignment(wrap_text=True, vertical="top")
        for column in wrap_columns:
            for row in ws.iter_rows(min_row=2, min_col=column, max_col=column):
                row[0].alignment = wrapped
    prompts = data["prompts"]
    snaps = data["snapshots"]
    pubs = data["publications"]
    start, end = data["start"], data["end"]
    cites_by_kind = {kind: citation_rows(pubs, snaps, start, end, kind) for kind in ("real", "manual", "simulated", "unknown")}
    match_map = {}
    for kind, rows in cites_by_kind.items():
        for pub in rows:
            for match in pub["matches"]:
                match_map.setdefault(match["sample_id"], []).append((match["kind"], pub["published_url"]))
    ws = sheet("样本明细", ["问题ID", "问题", "AI引擎", "日期(北京)", "样本ID", "样本来源", "提及", "引用URL", "精准引用我方URL", "宽松引用我方URL", "回答摘录"])
    engines = sorted({s.engine for s in snaps} | set(data.get("engines") or []))
    by_key = {}
    for s in snaps:
        d = local_time(s.captured_at).date()
        if start <= d <= end:
            by_key.setdefault((s.prompt_id, s.engine, d), []).append(s)
    days = (end - start).days + 1
    if len(prompts) * len(engines) * days + len(snaps) > 250000:
        raise ValueError("导出明细超过 25 万行，请缩短观察期或选择业务范围")
    cursor = start
    while cursor <= end:
        for p in prompts.values():
            for engine in engines:
                for s in by_key.get((p.id, engine, cursor), [None]):
                    matches = match_map.get(s.id, []) if s else []
                    values = [p.id, p.question, engine, cursor, s.id if s else None,
                              SOURCE_NAMES[provenance(s)] if s else "", "是" if s and s.mentions_brand else ("否" if s else ""),
                              "\n".join(s.cited_urls or []) if s else "",
                              "\n".join(u for k, u in matches if k == "exact"),
                              "\n".join(u for k, u in matches if k == "loose"),
                              (s.raw_text or "")[:500] if s else ""]
                    ws.append([safe_cell(v) for v in values])
                    ws.cell(ws.max_row, 4).number_format = "yyyy-mm-dd"
        cursor += timedelta(days=1)
    format_columns(ws, [8, 40, 16, 12, 8, 14, 8, 45, 45, 45, 60], (2, 8, 9, 10, 11))
    for gran, title in (("day", "日提及率"), ("week", "周提及率"), ("month", "月提及率")):
        ws = sheet(title, ["样本来源", "AI引擎", "周期", "提及数", "样本数", "提及率"])
        for kind in ("real", "manual", "simulated", "unknown"):
            for r in trend_rows(snaps, prompts, start, end, gran, kind, data.get("engines")):
                bucket = date.fromisoformat(r["bucket"]) if gran == "day" else r["bucket"]
                ws.append([safe_cell(v) for v in [SOURCE_NAMES[kind], r["engine"], bucket, r["mentions"], r["samples"], r["rate"]]])
                if gran == "day":
                    ws.cell(ws.max_row, 3).number_format = "yyyy-mm-dd"
                ws.cell(ws.max_row, 6).number_format = "0.0%"
        format_columns(ws, [14, 16, 12, 10, 10, 10])
    ws = sheet("URL引用汇总", ["样本来源", "发布ID", "任务ID", "渠道", "发布URL", "精准引用数", "宽松引用数", "命中样本ID"])
    for kind, rows in cites_by_kind.items():
        for r in rows:
            hit_ids = [m["sample_id"] for m in r["matches"]]
            hits = hit_ids[0] if len(hit_ids) == 1 else (",".join(map(str, hit_ids)) if hit_ids else None)
            ws.append([safe_cell(v) for v in [SOURCE_NAMES[kind], r["publication_id"], r["task_id"], r["channel"], r["published_url"], r["exact_count"], r["loose_count"], hits]])
    format_columns(ws, [14, 8, 8, 16, 45, 10, 10, 14], (5,))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _table(headers, rows):
    return "<table><thead><tr>" + "".join(f"<th>{escape(str(h))}</th>" for h in headers) + "</tr></thead><tbody>" + "".join(
        "<tr>" + "".join(f"<td>{escape(str(v)) if v is not None else '无数据'}</td>" for v in row) + "</tr>" for row in rows) + "</tbody></table>"


def _daily_line_chart(rows: list[dict], engine: str) -> str:
    """Draw sampled days only; missing days break the line instead of implying zero."""
    left, right, top, bottom = 48, 630, 14, 146
    count = len(rows)
    x_at = lambda i: left + (right - left) * i / max(count - 1, 1)
    y_at = lambda rate: bottom - (bottom - top) * rate
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" class="daily-trend" '
             f'width="650" height="184" viewBox="0 0 650 184" role="img" '
             f'aria-label="{escape(engine, quote=True)} 每日提及率">']
    for rate, label in ((1, "100%"), (.5, "50%"), (0, "0%")):
        y = y_at(rate)
        parts.append(f'<line x1="{left}" y1="{y:.1f}" x2="{right}" y2="{y:.1f}" stroke="#d7e2e9"/>')
        parts.append(f'<text x="2" y="{y + 4:.1f}" font-size="10">{label}</text>')
    path = []
    for i, row in enumerate(rows):
        if row["rate"] is None:
            continue
        x, y = x_at(i), y_at(row["rate"])
        path.append(f'{"M" if i == 0 or rows[i - 1]["rate"] is None else "L"}{x:.1f},{y:.1f}')
        parts.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="2.8" fill="#2875a9"/>')
    if path:
        parts.append(f'<path d="{" ".join(path)}" fill="none" stroke="#2875a9" stroke-width="2"/>')
    if rows:
        for i in sorted({0, (count - 1) // 2, count - 1}):
            parts.append(f'<text x="{x_at(i):.1f}" y="174" text-anchor="middle" font-size="10">{escape(rows[i]["bucket"])}</text>')
    parts.append('</svg>')
    return "".join(parts)


def _engine_trends(snaps, prompts, start: date, end: date, engines) -> str:
    daily = trend_rows(snaps, prompts, start, end, "day", engines=engines)
    weekly = trend_rows(snaps, prompts, start, end, "week", engines=engines)
    monthly = trend_rows(snaps, prompts, start, end, "month", engines=engines)
    names = sorted({r["engine"] for r in daily})
    blocks = []
    for engine in names:
        engine_daily = [r for r in daily if r["engine"] == engine]
        buckets = [("周", r) for r in weekly if r["engine"] == engine]
        buckets += [("月", r) for r in monthly if r["engine"] == engine]
        rows = [(kind, r["bucket"], r["mentions"], r["samples"],
                 f'{r["rate"]:.1%}' if r["rate"] is not None else "无数据") for kind, r in buckets]
        blocks.append(f'<div class="engine-trend"><h3>{escape(engine)}</h3>'
                      + _daily_line_chart(engine_daily, engine)
                      + _table(["粒度", "周期", "提及数", "样本数", "提及率"], rows) + '</div>')
    return "<p>主趋势：真实引擎采样。日维度明细见 Excel 导出。</p>" + "".join(blocks)


def build_report_html(data: dict, *, template: list[dict] | None = None, sample_label: str = "") -> str:
    start, end = data["start"], data["end"]
    prompts, snaps, pubs = data["prompts"], data["snapshots"], data["publications"]
    real = [s for s in snaps if provenance(s) == "real" and not getattr(prompts.get(s.prompt_id), "is_brand_probe", False)]
    manual = sum(provenance(s) == "manual" for s in snaps)
    simulated = sum(provenance(s) == "simulated" for s in snaps)
    unknown = sum(provenance(s) == "unknown" for s in snaps)
    mentioned = sum(bool(s.mentions_brand) for s in real)
    rate = f"{mentioned / len(real):.1%}" if real else "无数据"
    citations = citation_rows(pubs, snaps, start, end)
    matrix = []
    for prompt in prompts.values():
        for engine in sorted({s.engine for s in snaps} | set(data.get("engines") or [])):
            matching = [s for s in real if s.prompt_id == prompt.id and s.engine == engine]
            hits = sum(bool(s.mentions_brand) for s in matching)
            matrix.append((prompt.question, engine, hits, len(matching), f"{hits / len(matching):.1%}" if matching else "无数据"))
    blocks = {
        "overview": f"<p>真实样本 {len(real)} · 提及 {mentioned} · 提及率 {rate} · 问题 {len(prompts)} · 发布链接 {len(pubs)}</p>",
        "trends": _engine_trends(snaps, prompts, start, end, data.get("engines")),
        "citations": _table(["发布 URL", "精准匹配", "宽松匹配", "命中样本"], [(r["published_url"], r["exact_count"], r["loose_count"], ", ".join(str(m["sample_id"]) for m in r["matches"]) or "无数据") for r in citations]),
        "publications": _table(["渠道", "发布 URL", "发布时刻（北京）"], [(p.channel, p.published_url, local_time(p.published_at).strftime("%Y-%m-%d %H:%M") if p.published_at else "无数据") for p in pubs]),
        "provenance": f"<p>主指标仅统计真实引擎采样。人工录入 {manual} 条、模拟/演示 {simulated} 条、来源未知 {unknown} 条分别保留，不并入主指标。无样本显示无数据。</p>",
        "appendix": _table(["问题", "引擎", "提及数", "样本数", "提及率"], matrix[:100]) + ("<p>问题 × 引擎矩阵仅显示前 100 行；完整记录见 Excel。</p>" if len(matrix) > 100 else ""),
    }
    sections = "".join(f"<section><h2>{escape(item['title'])}</h2>{blocks[item['key']]}</section>" for item in (template or SECTION_DEFAULTS) if item["visible"])
    return f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><style>
      body {{ font-family: "Noto Sans CJK SC", "Noto Sans SC", "Source Han Sans SC", "WenQuanYi Micro Hei", "Microsoft YaHei", "PingFang SC", sans-serif; color:#19344a; font-size:11px; }}
      h1 {{ font-size:26px; }} h2 {{ font-size:16px; border-bottom:2px solid #2875a9; padding-bottom:5px; }}
      section {{ margin-top:20px; }} .engine-trend {{ break-inside:avoid; margin:14px 0; }} h3 {{ font-size:13px; margin:10px 0 3px; }} table {{ border-collapse:collapse; width:100%; table-layout:fixed; }} th,td {{ border:1px solid #cad6de; padding:5px; overflow-wrap:anywhere; }} th {{ background:#e8f2f7; }} svg {{ max-width:100%; }}
    </style></head><body><h1>{escape(sample_label + ' ' if sample_label else '')}GEO 优化报告</h1>
    <p>客户：{escape(str(data.get('client_name') or data['project_name']))} · 项目/业务：{escape(str(data['project_name']))}</p><p>观察期：{start.isoformat()} 至 {end.isoformat()}（北京时间）</p>
    <p>生成时间：{datetime.now(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m-%d %H:%M')} 北京时间</p>{sections}</body></html>'''
