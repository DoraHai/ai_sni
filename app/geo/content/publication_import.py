"""Parse and validate operator supplied publication links before the existing write gate."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from urllib.parse import urlsplit

from app.geo.content.report_data import normalize_report_url

MAX_BYTES = 2_000_000
MAX_ROWS = 500
HEADERS = {"task_id": {"task_id", "任务ID", "任务 id"}, "task_title": {"task_title", "任务标题", "标题", "问题", "question"},
           "channel": {"channel", "platform", "平台", "渠道"}, "url": {"url", "published_url", "发布链接", "发布URL", "链接"},
           "published_at": {"published_at", "publish_time", "发布时间", "发布时刻"}}


def clean(value) -> str:
    return re.sub(r"[\x00-\x1f\x7f]", "", str(value or "")).strip()


def parse_rows(*, filename: str = "", content: bytes | None = None, pasted: str = "") -> list[dict]:
    if content is not None and pasted:
        raise ValueError("文件和粘贴文本只能选择一种")
    if content is not None and len(content) > MAX_BYTES:
        raise ValueError("文件不能超过 2 MB")
    if content is None:
        content = pasted.encode("utf-8")
    if len(content) > MAX_BYTES:
        raise ValueError("文本不能超过 2 MB")
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = book.active
            raw = [[clean(c) for c in row] for row in sheet.iter_rows(max_row=MAX_ROWS + 2, values_only=True)]
            book.close()
        except Exception as exc:
            raise ValueError("无法读取 Excel 文件") from exc
    else:
        try:
            decoded = content.decode("utf-8-sig")
            raw = list(csv.reader(io.StringIO(decoded), delimiter="\t" if "\t" in decoded.splitlines()[0] else ","))
        except (UnicodeError, csv.Error, IndexError) as exc:
            raise ValueError("请使用 UTF-8 CSV 或制表符文本") from exc
    if not raw:
        raise ValueError("没有可导入的行")
    if len(raw) - 1 > MAX_ROWS:
        raise ValueError(f"最多导入 {MAX_ROWS} 行")
    names = [clean(x) for x in raw[0]]
    mapping = {key: next((i for i, name in enumerate(names) if name in aliases), None) for key, aliases in HEADERS.items()}
    if mapping["url"] is None or mapping["channel"] is None or (mapping["task_id"] is None and mapping["task_title"] is None):
        raise ValueError("缺少任务 ID/标题、平台/渠道或 URL 列")
    return [{"row_number": n, **{key: (clean(row[i]) if i is not None and i < len(row) else "") for key, i in mapping.items()}}
            for n, row in enumerate(raw[1:], 2) if any(clean(c) for c in row)]


def validate_rows(rows: list[dict], tasks: list, publications: list, business_id: int | None = None,
                  questions: dict[int, str] | None = None) -> list[dict]:
    by_id = {str(t.id): t for t in tasks}
    by_title = {}
    for t in tasks:
        by_title.setdefault(t.title, []).append(t)
        prompt_id = getattr(t, "prompt_id", None)
        if prompt_id is not None:
            by_title.setdefault(str(prompt_id), []).append(t)
            if questions and prompt_id in questions:
                by_title.setdefault(questions[prompt_id], []).append(t)
    existing = {(p.task_id, p.channel, normalize_report_url(p.published_url)) for p in publications}
    seen = set()
    out = []
    for row in rows:
        errors = []
        task_id = str(row.get("task_id") or "")
        task = by_id.get(task_id) if task_id else None
        if not task and not task_id and row.get("task_title"):
            choices = list({t.id: t for t in by_title.get(row["task_title"], [])}.values())
            if len(choices) == 1:
                task = choices[0]
            elif len(choices) > 1:
                errors.append("任务标题不唯一，请填写任务 ID")
        if not task:
            errors.append("任务不存在或不属于当前客户/项目")
        elif business_id is not None and task.business_id != business_id:
            errors.append("任务不属于当前项目")
        channel, url = row.get("channel", ""), row.get("url", "")
        if not channel or len(channel) > 32:
            errors.append("平台/渠道不能为空且不能超过 32 字")
        if len(url) > 2000:
            errors.append("URL 不能超过 2000 字")
        try:
            parsed = urlsplit(url)
            if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password or any(c.isspace() for c in url):
                errors.append("URL 必须是有效的 http(s) 链接")
        except ValueError:
            errors.append("URL 必须是有效的 http(s) 链接")
        published_at = None
        if row.get("published_at"):
            try:
                published_at = datetime.fromisoformat(row["published_at"].replace("Z", "+00:00"))
            except ValueError:
                errors.append("发布时间格式无效，请使用 YYYY-MM-DD HH:MM:SS")
        key = (task.id if task else None, channel, normalize_report_url(url))
        if key in seen or key in existing:
            errors.append("发布链接重复")
        seen.add(key)
        out.append({**row, "task_id": task.id if task else task_id, "published_at": published_at.isoformat() if published_at else "",
                    "status": "有效" if not errors else "错误", "errors": errors})
    return out
