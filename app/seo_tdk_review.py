"""Pure, stored-data-only TDK review context and Word/HTML layouts."""
from __future__ import annotations

import base64
from datetime import datetime, timezone
from html import escape
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt
from PIL import Image

from app.seo_publication_export import beijing_text

DEFAULT_TDK_REVIEW_SECTIONS = (
    {"key": "header", "enabled": True, "title": "页面信息"},
    {"key": "screenshot", "enabled": True, "title": "页面截图"},
    {"key": "tdk_compare", "enabled": True, "title": "TDK 对比"},
    {"key": "keyword_layout", "enabled": True, "title": "关键词布局"},
    {"key": "internal_links", "enabled": True, "title": "内链数据"},
    {"key": "client_comment", "enabled": True, "title": "客户意见"},
)
DEFAULT_TDK_REVIEW_COLUMNS = {"char_counts": True, "rationale": True}


def validate_tdk_review_template(sections, columns=None):
    if not isinstance(sections, list) or len(sections) != len(DEFAULT_TDK_REVIEW_SECTIONS):
        raise ValueError("审核稿章节须包含全部六个章节")
    keys = [item.get("key") if isinstance(item, dict) else None for item in sections]
    allowed = {item["key"] for item in DEFAULT_TDK_REVIEW_SECTIONS}
    if set(keys) != allowed or len(set(keys)) != len(keys):
        raise ValueError("审核稿章节无效或重复")
    normalized = []
    for item in sections:
        title = item.get("title")
        if not isinstance(item.get("enabled"), bool) or not isinstance(title, str) or not 1 <= len(title.strip()) <= 40:
            raise ValueError("章节显示状态或标题无效")
        normalized.append({"key": item["key"], "enabled": item["enabled"], "title": title.strip()})
    if not any(item["enabled"] for item in normalized):
        raise ValueError("至少显示一个章节")
    columns = DEFAULT_TDK_REVIEW_COLUMNS.copy() if columns is None else columns
    if not isinstance(columns, dict) or set(columns) != set(DEFAULT_TDK_REVIEW_COLUMNS) or any(type(value) is not bool for value in columns.values()):
        raise ValueError("TDK 表格列设置无效")
    return {"sections": normalized, "columns": dict(columns)}


def _get(obj, key, default=None):
    return obj.get(key, default) if isinstance(obj, dict) else getattr(obj, key, default)


def xml_text(value):
    """Word XML 1.0 cannot contain most C0 controls or surrogate codepoints."""
    return "".join(c for c in str(value if value is not None else "")
                   if ord(c) in (9, 10, 13) or 32 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)


def _display(value, empty="空"):
    value = xml_text(value).strip()
    return value or empty


def build_tdk_review_context(*, site_name, pages, batch_id, generated_at=None,
                             keywords=(), links=(), captures=(), images=None,
                             capture_notes=None, sample_label=None, ai_suggestions=None):
    """Accept already scoped rows; derive only literal keyword substring matches."""
    generated_at = generated_at or datetime.now(timezone.utc)
    keyword_rows = list(keywords)
    links = list(links)
    captures = list(captures)
    pages = list(pages)
    page_by_id = {_get(page, "id"): page for page in pages}
    rendered = []
    for page in pages:
        page_id, url = _get(page, "id"), _get(page, "url")
        ai = (ai_suggestions or {}).get(page_id)
        capture = next((row for row in sorted(captures, key=lambda row: (_get(row, "captured_at") or datetime.min.replace(tzinfo=timezone.utc), _get(row, "id") or 0), reverse=True)
                        if _get(row, "relation_id") == page_id and _get(row, "status") == "succeeded" and _get(row, "storage_key")), None)
        tdk = []
        for marker, label, current, suggested in (
            ("①", "Title", _get(page, "title"), _get(page, "title_suggestion")),
            ("②", "Description", _get(page, "meta_description"), _get(page, "description_suggestion")),
            ("③", "Keywords", _get(page, "meta_keywords"), None),
        ):
            field = label.lower()
            ai_status = _get(ai, f"{field}_status")
            ai_confirmed = ai_status in {"confirmed", "modified"}
            if ai_confirmed:
                suggested = _get(ai, f"{field}_final_value")
            source = ("AI建议（人工修改后确认）" if ai_status == "modified" else "AI建议（已人工确认）") if ai_confirmed else ("系统建议" if suggested else "")
            tdk.append({"marker": marker, "label": label, "current": _display(current),
                        "suggested": _display(suggested, "暂无建议"), "suggestion_source": source,
                        "counts": f"{len(xml_text(current or ''))} / {len(xml_text(suggested or ''))}",
                        "rationale": _display(_get(ai, "reason"), "") if ai_confirmed else "", "client_comment": ""})
        mapped = []
        seen = set()
        for keyword in keyword_rows:
            if _get(keyword, "id") == _get(page, "target_keyword_id") or (_get(keyword, "landing_page") and _get(keyword, "landing_page") == url):
                word = _display(_get(keyword, "keyword"), "")
                if not word or word.casefold() in seen:
                    continue
                seen.add(word.casefold())
                has = lambda value: bool(value and word.casefold() in xml_text(value).casefold())
                mapped.append({"keyword": word, "title": has(_get(page, "title")),
                               "description": has(_get(page, "meta_description")), "h1": has(_get(page, "h1")),
                               "suggested_title": has(_get(page, "title_suggestion")),
                               "suggested_description": has(_get(page, "description_suggestion"))})
        edges = []
        for link in links:
            source_id, target_id = _get(link, "source_page_id"), _get(link, "target_page_id")
            if page_id in (source_id, target_id) and source_id in page_by_id and target_id in page_by_id:
                edges.append({"source": _get(page_by_id[source_id], "url"),
                              "target": _get(page_by_id[target_id], "url"),
                              "anchor": _display(_get(link, "anchor_text"), "空"), "source_label": "已发现内链"})
        for link in _get(ai, "internal_link_suggestions", []) or []:
            if _get(link, "status") in {"confirmed", "modified"} and _get(link, "target_page_id") in page_by_id and _get(link, "target_page_id") != page_id:
                edges.append({"source": url, "target": _get(page_by_id[_get(link, "target_page_id")], "url"),
                              "anchor": _display(_get(link, "final_anchor") or _get(link, "anchor"), "空"),
                              "source_label": "AI内链建议（已人工确认）"})
        rendered.append({"id": page_id, "url": url, "title": _display(_get(page, "title")),
                         "tdk": tdk, "keywords": mapped, "links": edges,
                         "image_key": _get(capture, "storage_key") if capture else None,
                         "capture_note": (capture_notes or {}).get(page_id, "暂无截图" if not capture else ""),
                         "status": _get(page, "status")})
    return {"site": _display(site_name), "batch_id": batch_id,
            "generated_at": beijing_text(generated_at) + " 北京时间",
            "pages": rendered, "images": images or {}, "sample_label": sample_label}


def _image_jpeg(data, *, max_width=1400):
    with Image.open(BytesIO(data)) as source:
        if source.width * source.height > 16_000_000:
            raise ValueError("截图图片过大")
        image = source.convert("RGB")
        if image.width > max_width:
            image.thumbnail((max_width, 3000))
        output = BytesIO()
        image.save(output, "JPEG", quality=82, optimize=True)
        return output.getvalue()


def _sections(template):
    if template is None:
        return validate_tdk_review_template(list(DEFAULT_TDK_REVIEW_SECTIONS))
    if isinstance(template, dict):
        return validate_tdk_review_template(template.get("sections"), template.get("columns"))
    return validate_tdk_review_template(template.sections, template.columns)


def render_tdk_review_docx(context, template=None, image_loader=None) -> bytes:
    settings = _sections(template)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Microsoft YaHei", Pt(9)
    document.sections[0].header.paragraphs[0].text = xml_text(context.get("sample_label") or "")
    document.sections[0].top_margin = Cm(1.8)
    document.sections[0].bottom_margin = Cm(1.6)
    for number, page in enumerate(context["pages"]):
        if number:
            document.add_page_break()
        for section in settings["sections"]:
            if not section["enabled"]:
                continue
            key = section["key"]
            document.add_heading(xml_text(section["title"]), level=1 if key == "header" else 2)
            if key == "header":
                for label, value in (("网站", context["site"]), ("页面 URL", page["url"]),
                                     ("页面标题", page["title"]), ("导出时间", context["generated_at"]),
                                     ("批次 ID", context["batch_id"])):
                    document.add_paragraph(f"{label}：{xml_text(value)}")
            elif key == "screenshot":
                try:
                    data = (image_loader(page["image_key"]) if image_loader else context.get("images", {}).get(page["image_key"])) if page["image_key"] else None
                    if data:
                        document.add_picture(BytesIO(_image_jpeg(data)), width=Cm(16))
                    else:
                        document.add_paragraph(xml_text(page["capture_note"] or "暂无截图"))
                except (OSError, ValueError):
                    document.add_paragraph("暂无截图（图片不可用）")
                document.add_paragraph("① Title　② Description　③ Keywords（对应下表字段；截图不标注不可见的 meta 标签）")
            elif key == "tdk_compare":
                columns = ["序号", "字段", "当前值", "建议值"]
                if settings["columns"]["char_counts"]:
                    columns.append("字数（当前 / 建议）")
                if settings["columns"]["rationale"]:
                    columns.append("说明 / 批注")
                columns.append("客户意见")
                table = document.add_table(rows=1, cols=len(columns))
                table.style = "Table Grid"
                for cell, value in zip(table.rows[0].cells, columns):
                    cell.text = value
                for item in page["tdk"]:
                    values = [item["marker"], item["label"], item["current"],
                              (item.get("suggestion_source", "") + "：" if item.get("suggestion_source") else "") + item["suggested"]]
                    if settings["columns"]["char_counts"]:
                        values.append(item["counts"])
                    if settings["columns"]["rationale"]:
                        values.append(item["rationale"])
                    values.append("")
                    for cell, value in zip(table.add_row().cells, values):
                        cell.text = xml_text(value)
            elif key == "keyword_layout":
                if not page["keywords"]:
                    document.add_paragraph("系统中暂无该页面的关键词映射")
                else:
                    table = document.add_table(rows=1, cols=6)
                    table.style = "Table Grid"
                    for cell, value in zip(table.rows[0].cells, ("关键词", "当前 Title", "当前 Description", "当前 H1", "建议 Title", "建议 Description")):
                        cell.text = value
                    for item in page["keywords"]:
                        for cell, value in zip(table.add_row().cells, (item["keyword"], *("是" if item[key] else "否" for key in ("title", "description", "h1", "suggested_title", "suggested_description")))):
                            cell.text = xml_text(value)
            elif key == "internal_links":
                if not page["links"]:
                    document.add_paragraph("暂无内链建议数据")
                else:
                    for link in page["links"]:
                        document.add_paragraph(xml_text(f"{link['source_label']}\n来源：{link['source']}\n目标：{link['target']}\n锚文本：{link['anchor']}"))
            elif key == "client_comment":
                document.add_paragraph("审核结论：□ 同意　□ 修改　□ 驳回")
                document.add_paragraph("意见：________________________________________________________")
                document.add_paragraph("____________________________________________________________")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_tdk_review_html(context, template=None):
    settings = _sections(template)
    e = lambda value: escape(xml_text(value), quote=True)
    body = []
    for page in context["pages"]:
        body.append('<article class="page">')
        for section in settings["sections"]:
            if not section["enabled"]:
                continue
            key = section["key"]
            body.append(f'<h2>{e(section["title"])}</h2>')
            if key == "header":
                for label, value in (("网站", context["site"]), ("页面 URL", page["url"]), ("页面标题", page["title"]),
                                     ("导出时间", context["generated_at"]), ("批次 ID", context["batch_id"])):
                    body.append(f'<p><b>{label}：</b>{e(value)}</p>')
            elif key == "screenshot":
                data = context.get("images", {}).get(page["image_key"]) if page["image_key"] else None
                try:
                    if data:
                        encoded = base64.b64encode(_image_jpeg(data)).decode("ascii")
                        body.append(f'<img class="shot" src="data:image/jpeg;base64,{encoded}">')
                    else:
                        body.append(f'<p>{e(page["capture_note"] or "暂无截图")}</p>')
                except (OSError, ValueError):
                    body.append('<p>暂无截图（图片不可用）</p>')
                body.append('<p>① Title　② Description　③ Keywords（对应下表字段；截图不标注不可见的 meta 标签）</p>')
            elif key == "tdk_compare":
                cols = [("marker", "序号"), ("label", "字段"), ("current", "当前值"), ("suggested", "建议值")]
                if settings["columns"]["char_counts"]:
                    cols.append(("counts", "字数（当前 / 建议）"))
                if settings["columns"]["rationale"]:
                    cols.append(("rationale", "说明 / 批注"))
                cols.append(("client_comment", "客户意见"))
                body.append('<table><tr>' + ''.join(f'<th>{e(title)}</th>' for _, title in cols) + '</tr>')
                for item in page["tdk"]:
                    cells = []
                    for field, _ in cols:
                        value = item[field]
                        if field == "suggested" and item.get("suggestion_source"):
                            value = item["suggestion_source"] + "：" + value
                        cells.append(f'<td>{e(value)}</td>')
                    body.append('<tr>' + ''.join(cells) + '</tr>')
                body.append('</table>')
            elif key == "keyword_layout":
                if not page["keywords"]:
                    body.append('<p>系统中暂无该页面的关键词映射</p>')
                else:
                    fields = ("title", "description", "h1", "suggested_title", "suggested_description")
                    body.append('<table><tr><th>关键词</th><th>当前 Title</th><th>当前 Description</th><th>当前 H1</th><th>建议 Title</th><th>建议 Description</th></tr>')
                    for item in page["keywords"]:
                        body.append('<tr><td>' + e(item["keyword"]) + '</td>' + ''.join(f'<td>{"是" if item[field] else "否"}</td>' for field in fields) + '</tr>')
                    body.append('</table>')
            elif key == "internal_links":
                if not page["links"]:
                    body.append('<p>暂无内链建议数据</p>')
                else:
                    for link in page["links"]:
                        body.append(f'<p>{e(link["source_label"])}<br>来源：{e(link["source"])}<br>目标：{e(link["target"])}<br>锚文本：{e(link["anchor"])}</p>')
            elif key == "client_comment":
                body.append('<p>审核结论：□ 同意　□ 修改　□ 驳回</p><p>意见：____________________________________________</p><p>____________________________________________</p>')
        body.append('</article>')
    style = "@page{size:A4;margin:18mm 14mm}body{font-family:'Microsoft YaHei','Noto Sans CJK SC',sans-serif;color:#203247;font-size:9pt}.page{break-before:page}.page:first-child{break-before:auto}h2{color:#194c75;font-size:13pt;margin:11px 0 5px}p{margin:5px 0;overflow-wrap:anywhere}table{border-collapse:collapse;width:100%;table-layout:fixed;font-size:8pt}th,td{border:1px solid #b8c9d8;padding:4px;vertical-align:top;overflow-wrap:anywhere}th{background:#eaf1f7}.shot{display:block;max-width:16cm;max-height:9cm;object-fit:contain}.page{page-break-inside:avoid}"
    return '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><style>' + style + '</style></head><body>' + ''.join(body) + '</body></html>'
