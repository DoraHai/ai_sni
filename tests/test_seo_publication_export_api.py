"""Publication export API with SQLite metadata and no external services."""

from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.schema import CreateTable

from app.api import seo_page_captures as api
from app.api.seo import router as seo_router, require_seo_module_access
from app.models.module_workspace import SeoSite
from app.models.seo import SeoContentAsset, SeoContentPublication, SeoKeywordAsset
from app.models.seo_page_capture import SeoPageCapture
from app.security.auth import AuthContext
from app.seo_demo_source import require_seo_auth


@compiles(JSONB, "sqlite")
def _jsonb_sqlite(_type, _compiler, **_kwargs):
    return "JSON"


class Adapter:
    def __init__(self, db):
        self.db = db

    async def get(self, model, ident):
        return self.db.get(model, ident)

    async def execute(self, statement):
        return self.db.execute(statement)

    async def scalars(self, statement):
        return self.db.scalars(statement)


def _client(monkeypatch, actor):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    with engine.begin() as connection:
        for model in (SeoSite, SeoContentAsset, SeoContentPublication, SeoKeywordAsset, SeoPageCapture):
            connection.execute(CreateTable(model.__table__, include_foreign_key_constraints=[]))
    with Session(engine) as db:
        db.add_all([
            SeoSite(id=2, tenant_id=4, tenant_module_id=1, name="测试站", domain="example.com", canonical_domain="example.com", status="active"),
            SeoSite(id=3, tenant_id=5, tenant_module_id=2, name="其他站", domain="other.com", canonical_domain="other.com", status="active"),
            SeoContentAsset(id=11, tenant_id=4, site_id=2, title="原文", keyword_id=31, keyword_ids=[31]),
            SeoContentAsset(id=12, tenant_id=5, site_id=3, title="跨租户文章"),
            SeoKeywordAsset(id=31, tenant_id=4, site_id=2, keyword="关键词"),
            SeoContentPublication(id=21, tenant_id=4, content_asset_id=11, platform_code="zhihu", platform_name="知乎", status="published", page_url="https://example.com/a", published_at=datetime(2026, 9, 30, 15, 59)),
            SeoContentPublication(id=22, tenant_id=4, content_asset_id=11, platform_code="web", platform_name="官网", status="published", page_url=None, published_at=datetime(2026, 9, 30, 16, 30)),
            SeoContentPublication(id=23, tenant_id=4, content_asset_id=11, platform_code="web", platform_name="官网", status="draft_created", published_at=datetime(2026, 9, 12)),
            SeoContentPublication(id=24, tenant_id=5, content_asset_id=12, platform_code="web", platform_name="其他", status="published", published_at=datetime(2026, 9, 12)),
            SeoPageCapture(id=1, tenant_id=4, site_id=2, relation_type="publication", relation_id=21, source_url="https://example.com/a", status="failed", error_code="timeout", captured_at=datetime(2026, 9, 29, tzinfo=timezone.utc), viewport_width=800, viewport_height=600, redirect_chain=[], warnings={}),
        ])
        db.commit()
    app = FastAPI()
    app.include_router(seo_router)

    async def auth():
        if actor == "anonymous":
            raise HTTPException(401, "未登录")
        return actor

    async def session():
        with Session(engine) as db:
            yield Adapter(db)

    app.dependency_overrides[require_seo_auth] = auth
    app.dependency_overrides[require_seo_module_access] = auth
    app.dependency_overrides[api.get_session] = session
    monkeypatch.setattr(api, "get_settings", lambda: SimpleNamespace(seo_page_capture_storage_dir="C:/missing", seo_page_capture_max_pixels=16_000_000))
    return TestClient(app)


def _actor(tenant=4, permission="view"):
    return AuthContext(9, "tester", "client", tenant, {"seo.site": permission})


def test_export_permissions_month_filter_and_headers(monkeypatch):
    path = "/api/v1/seo/site/publications/export"
    params = {"tenant_id": 4, "site_id": 2, "month": "2026-09"}
    assert _client(monkeypatch, "anonymous").get(path, params=params).status_code == 401
    assert _client(monkeypatch, _actor(5)).get(path, params=params).status_code == 403
    assert _client(monkeypatch, _actor(permission="none")).get(path, params=params).status_code == 403
    assert _client(monkeypatch, _actor()).get(path, params={**params, "site_id": 3}).status_code == 404
    response = _client(monkeypatch, _actor()).get(path, params=params)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "private, no-store"
    workbook = load_workbook(BytesIO(response.content))
    detail = workbook["发布明细"]
    assert detail.max_row == 2
    assert detail["B2"].value == "知乎"
    assert detail["C2"].value == "原文"
    assert detail["D2"].value == "关键词"
    assert detail["G2"].value == "失败：截图超时"
    assert detail["F2"].value == "2026-09-30 23:59"
    october = _client(monkeypatch, _actor()).get(path, params={**params, "month": "2026-10"})
    october_sheet = load_workbook(BytesIO(october.content))["发布明细"]
    assert october_sheet.max_row == 2
    assert october_sheet["B2"].value == "官网"
    assert october_sheet["E2"].value is None
    assert _client(monkeypatch, _actor()).get(path, params={**params, "month": "2026-13"}).status_code == 422
