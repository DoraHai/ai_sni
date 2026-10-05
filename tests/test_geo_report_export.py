import asyncio
import re
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.geo.report_export import build_report_html, build_xlsx, report_template
from app.geo.report_pdf import render_report_pdf
from app.geo.content.report_routes import mention_trends, _project
from app.security.auth import _required


def fixture_data():
    prompt = SimpleNamespace(id=1, question='=danger', is_brand_probe=False)
    real = SimpleNamespace(id=1, prompt_id=1, engine='E', captured_at=datetime(2026, 9, 1),
                           sample_mode='openai_compat', patrol_run_id=1, simulated=False, note='', mentions_brand=True,
                           cited_urls=['https://example.com/a'], raw_text='<script>alert(1)</script>')
    manual = SimpleNamespace(**{**vars(real), 'id': 2, 'sample_mode': 'manual', 'mentions_brand': False})
    sim = SimpleNamespace(**{**vars(real), 'id': 3, 'simulated': True, 'sample_mode': 'mock_persona'})
    pub = SimpleNamespace(id=2, task_id=5, channel='web', published_url='https://example.com/a', published_at=None)
    return dict(project_name='<client>', start=date(2026, 9, 1), end=date(2026, 9, 1),
                prompts={1: prompt}, snapshots=[real, manual, sim], publications=[pub], engines=['E'])


def test_workbook_has_all_sources_and_injection_guard_and_empty_rate():
    label = '【示例数据】仅供格式预览'
    wb = load_workbook(BytesIO(build_xlsx(fixture_data(), sample_label=label)))
    assert {'样本明细', '日提及率', '周提及率', '月提及率', 'URL引用汇总'} <= set(wb.sheetnames)
    for ws in wb:
        assert ws['A1'].value == label
    detail = wb['样本明细']
    assert detail['A3'].value == 1 and detail['A3'].data_type == 'n'
    assert detail['D3'].value == datetime(2026, 9, 1) and detail['D3'].number_format == 'yyyy-mm-dd'
    assert detail['E3'].value == 1 and detail['E3'].data_type == 'n'
    assert detail['B3'].value == "'=danger"
    assert {detail.cell(i, 6).value for i in (3, 4, 5)} == {'真实引擎采样', '人工录入', '模拟/演示'}
    assert detail['K3'].value.startswith('<script>')
    rates = wb['日提及率']
    assert rates['F3'].value == 1
    assert rates['F4'].value == 0
    assert rates['D3'].value == 1 and rates['D3'].data_type == 'n'
    assert rates['E3'].value == 1 and rates['E3'].data_type == 'n'
    assert rates['F3'].data_type == 'n' and rates['F3'].number_format == '0.0%'
    assert rates['C3'].value == datetime(2026, 9, 1) and rates['C3'].number_format == 'yyyy-mm-dd'
    citation = wb['URL引用汇总']
    assert citation['B3'].value == 2 and citation['B3'].data_type == 'n'
    assert citation['C3'].value == 5 and citation['C3'].data_type == 'n'
    assert citation['F3'].value == 1 and citation['F3'].data_type == 'n'
    assert citation['H3'].value == 1 and citation['H3'].data_type == 'n'
    assert citation['H4'].value == 2 and citation['H4'].data_type == 'n'


def test_empty_rate_is_blank_number_cell():
    data = fixture_data()
    data['snapshots'] = []
    wb = load_workbook(BytesIO(build_xlsx(data)))
    assert wb['日提及率']['F3'].value is None
    assert wb['日提及率']['F3'].number_format == '0.0%'
    assert wb['样本明细']['E3'].value is None


def test_workbook_column_layout_and_header_controls():
    wb = load_workbook(BytesIO(build_xlsx(fixture_data(), sample_label='【示例数据】')))
    layouts = {
        '样本明细': ([8, 40, 16, 12, 8, 14, 8, 45, 45, 45, 60], {'B', 'H', 'I', 'J', 'K'}),
        '日提及率': ([14, 16, 12, 10, 10, 10], set()),
        '周提及率': ([14, 16, 12, 10, 10, 10], set()),
        '月提及率': ([14, 16, 12, 10, 10, 10], set()),
        'URL引用汇总': ([14, 8, 8, 16, 45, 10, 10, 14], {'E'}),
    }
    for name, (widths, wrapped) in layouts.items():
        ws = wb[name]
        assert ws.freeze_panes == 'A3'
        assert ws.auto_filter.ref == f'A2:{ws.cell(2, len(widths)).column_letter}2'
        for column, width in enumerate(widths, 1):
            assert ws.column_dimensions[ws.cell(2, column).column_letter].width == width
        for letter in wrapped:
            assert ws[f'{letter}3'].alignment.wrap_text is True


def test_html_escapes_and_template_defaults():
    html = build_report_html(fixture_data(), sample_label='示例数据')
    assert '&lt;client&gt;' in html and '<script>' not in html
    assert '人工录入 1' in html and '模拟/演示 1' in html
    assert len(report_template(None)) == 6
    assert re.search(r'生成时间：\d{4}-\d{2}-\d{2} \d{2}:\d{2} 北京时间', html)
    assert 'T01:' not in html


def test_daily_svg_gaps_and_weekly_monthly_sums():
    data = fixture_data()
    real = data['snapshots'][0]
    data['end'] = date(2026, 9, 8)
    data['snapshots'] = [
        real,
        SimpleNamespace(**{**vars(real), 'id': 4, 'mentions_brand': False}),
        SimpleNamespace(**{**vars(real), 'id': 5, 'captured_at': datetime(2026, 9, 3)}),
    ]
    html = build_report_html(data)
    assert '<svg' in html and 'class="daily-trend"' in html
    assert re.search(r'<path d="M[^"<]* M[^"<]*" fill="none"', html)
    assert '<th>北京日期</th>' not in html
    assert '日维度明细见 Excel 导出' in html
    assert '<td>周</td><td>2026-W36</td><td>2</td><td>3</td><td>66.7%</td>' in html
    assert '<td>周</td><td>2026-W37</td><td>0</td><td>0</td><td>无数据</td>' in html
    assert '<td>月</td><td>2026-09</td><td>2</td><td>3</td><td>66.7%</td>' in html
    wb = load_workbook(BytesIO(build_xlsx(data)))
    weekly = wb['周提及率']
    assert [weekly.cell(3, col).value for col in (4, 5)] == [2, 3]
    assert weekly['F3'].value == pytest.approx(2 / 3)
    assert weekly['F3'].data_type == 'n' and weekly['F3'].number_format == '0.0%'


def test_report_routes_require_geo_content_edit_for_mutations():
    assert _required('/api/v1/geo/reports/export.pdf', 'GET') == ({'geo.content'}, False)
    assert _required('/api/v1/geo/reports/publication-links/confirm', 'POST') == ({'geo.content'}, True)
    assert _required('/api/v1/geo/reports/template', 'PUT') == ({'geo.content'}, True)


def test_pdf_fake_browser_and_missing_runtime_503():
    asyncio.run(_pdf_fake_browser_and_missing_runtime_503())


async def _pdf_fake_browser_and_missing_runtime_503():
    seen = {}
    class Page:
        async def route(self, pattern, handler): seen['route'] = pattern; seen['handler'] = handler
        async def set_content(self, html, **kw): seen['html'] = html
        async def pdf(self, **kw): seen['pdf'] = kw; return b'%PDF-fake'
    class Context:
        async def new_page(self): return Page()
        async def close(self): pass
    class Browser:
        async def new_context(self, **kw): seen['context'] = kw; return Context()
        async def close(self): pass
    class Chromium:
        async def launch(self, **kw): seen['launch'] = kw; return Browser()
    class PW:
        chromium = Chromium()
        async def __aenter__(self): return self
        async def __aexit__(self, *a): pass
    settings = SimpleNamespace(geo_report_browser_executable_path='C:/chrome.exe', geo_report_browser_channel='')
    assert await render_report_pdf('<p>safe</p>', sample_label='示例数据', playwright_factory=PW, settings=settings) == b'%PDF-fake'
    assert seen['context']['offline'] is True and seen['context']['java_script_enabled'] is False
    assert '示例数据' in seen['pdf']['footer_template']
    assert 'pageNumber' in seen['pdf']['footer_template'] and 'totalPages' in seen['pdf']['footer_template']
    class Route:
        def __init__(self, url): self.request = SimpleNamespace(url=url); self.action = None
        async def abort(self): self.action = 'abort'
        async def continue_(self): self.action = 'continue'
    external, inline = Route('https://example.com/x'), Route('data:image/png;base64,AA==')
    await seen['handler'](external)
    await seen['handler'](inline)
    assert (external.action, inline.action) == ('abort', 'continue')
    class Broken(Chromium):
        async def launch(self, **kw): raise FileNotFoundError('missing')
    class BrokenPW(PW):
        chromium = Broken()
    with pytest.raises(HTTPException) as err:
        await render_report_pdf('<p/>', playwright_factory=BrokenPW, settings=settings)
    assert err.value.status_code == 503


def test_route_checks_tenant_before_reading_any_data():
    asyncio.run(_route_checks_tenant_before_reading_any_data())


async def _route_checks_tenant_before_reading_any_data():
    class Context:
        def ensure_tenant(self, tenant_id): raise HTTPException(403, '租户无权限')
    with pytest.raises(HTTPException) as err:
        await mention_trends(tenant_id=999, business_id=None, from_=date(2026, 9, 1), to=date(2026, 9, 2),
                             granularity='day', provenance='real', ctx=Context(), session=None)
    assert err.value.status_code == 403
    class MissingProjectSession:
        async def scalar(self, statement): return None
    with pytest.raises(HTTPException) as missing:
        await _project(MissingProjectSession(), 1, 999)
    assert missing.value.status_code == 404
