"""Stored-only TDK context, Word and HTML layout contracts."""
from io import BytesIO
from types import SimpleNamespace as Row

from docx import Document
import pytest

from app.seo_tdk_review import (DEFAULT_TDK_REVIEW_SECTIONS, build_tdk_review_context,
    render_tdk_review_docx, render_tdk_review_html, validate_tdk_review_template)
from scripts.generate_tdk_review_sample import sample_context


def _text(doc):
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for table in doc.tables for row in table.rows for c in row.cells])


def test_context_uses_only_stored_suggestions_and_keyword_mapping():
    page = Row(id=1, url="https://example.invalid/a", title="A Keyword", meta_description="some content",
        meta_keywords=None, h1="KEYWORD", title_suggestion="new keyword", description_suggestion=None,
        target_keyword_id=5, status="proposed")
    context = build_tdk_review_context(site_name="site", pages=[page], batch_id=1,
        keywords=[Row(id=5, keyword="keyword", landing_page=None)], links=[])
    item = context["pages"][0]
    assert item["tdk"][1]["suggested"] == "暂无建议"
    assert item["tdk"][2]["suggested"] == "暂无建议"
    assert item["keywords"][0]["title"] and item["keywords"][0]["h1"]
    assert item["keywords"][0]["suggested_title"] and not item["keywords"][0]["description"]
    assert not item["links"]
    assert "暂无内链建议数据" in render_tdk_review_html(context)


def test_template_validates_keys_order_and_toggles():
    sections = list(DEFAULT_TDK_REVIEW_SECTIONS)[::-1]
    assert validate_tdk_review_template(sections)["sections"][0]["key"] == "client_comment"
    with pytest.raises(ValueError):
        validate_tdk_review_template(sections[:-1] + [sections[0]])
    with pytest.raises(ValueError):
        validate_tdk_review_template(sections, {"char_counts": "yes", "rationale": True})


def test_docx_opens_contains_image_annotations_and_header_label():
    context = sample_context()
    assert len(context["pages"]) == 3
    assert context["pages"][0]["links"] and not context["pages"][1]["links"]
    assert context["pages"][1]["tdk"][0]["suggested"] != "暂无建议"
    assert context["pages"][2]["tdk"][0]["suggested"] == "暂无建议"
    data = render_tdk_review_docx(context)
    doc = Document(BytesIO(data))
    text = _text(doc)
    assert "①" in text and "②" in text and "③" in text
    assert "暂无截图" in text and "暂无建议" in text
    assert "【示例数据】演示站点" in text
    assert "示例数据" in doc.sections[0].header.paragraphs[0].text
    assert len(doc.inline_shapes) == 1


def test_html_escapes_script_and_word_strips_xml_invalid_controls():
    context = sample_context()
    context["pages"][0]["title"] = "<script>alert(1)</script>\x01"
    context["pages"][0]["tdk"][0]["current"] = "unsafe\x02"
    html = render_tdk_review_html(context)
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "\x01" not in html and "\x02" not in html
    doc = Document(BytesIO(render_tdk_review_docx(context)))
    assert "\x01" not in _text(doc) and "\x02" not in _text(doc)


def test_ai_export_requires_human_confirmation_for_fields_and_links():
    context = sample_context()
    page = Row(id=51, url="https://example.invalid/p", title="当前标题", meta_description="当前描述",
        meta_keywords=None, h1=None, title_suggestion="系统标题", description_suggestion=None,
        target_keyword_id=None, status="pending")
    target = Row(id=52, url="https://example.invalid/target", title="目标页", meta_description=None,
        meta_keywords=None, h1=None, title_suggestion=None, description_suggestion=None,
        target_keyword_id=None, status="pending")
    ai = Row(title_status="confirmed", title_final_value="确认标题", description_status="ai_draft",
        description_ai_value="绝不能导出的草稿", keywords_status="rejected", keywords_ai_value="绝不能导出的驳回词",
        reason="人工确认理由", internal_link_suggestions=[
            {"target_page_id": 52, "anchor": "已确认锚文本", "status": "confirmed", "final_anchor": "已确认锚文本"},
            {"target_page_id": 52, "anchor": "绝不能导出的草稿锚文本", "status": "ai_draft"}])
    context = build_tdk_review_context(site_name="测试站", pages=[page, target], batch_id=1,
        ai_suggestions={51: ai})
    context["pages"] = context["pages"][:1]
    text = _text(Document(BytesIO(render_tdk_review_docx(context))))
    html = render_tdk_review_html(context)
    for content in (text, html):
        assert "AI建议（已人工确认）" in content and "确认标题" in content
        assert "人工确认理由" in content and "AI内链建议（已人工确认）" in content
        assert "系统标题" not in content and "绝不能导出的" not in content
        assert "已确认锚文本" in content and "暂无建议" in content
