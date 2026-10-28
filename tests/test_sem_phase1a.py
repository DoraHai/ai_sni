"""全程本地合成；不导入main/config，不读取.env，不使用HTTP/OCR/模型/百度。"""
import copy
import io
import json
import socket
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.sem_phase0.contracts import LandingSubmission
from app.sem_phase0.store import AdvisorPrincipal, RevisionConflict, ScopeDenied, SubmissionStore
from app.sem_phase1a.contracts import (ApplyRequest, ApprovalRequest, CreateRequest, RulingRequest,
                                      PlanRow, WARNING)
from app.sem_phase1a.excel import MAX_XLSX_BYTES, parse_workbook
from app.sem_phase1a.planning import duplicates
from app.sem_phase1a.router import create_router
from app.sem_phase1a.service import ShadowService, digest, enable_plan, export_execution
from app.sem_phase1a.sources import FIXTURES, fixture_adapters
from app.sem_phase1a.principals import SessionIdentity, principal_dependency
from sem_phase1a_auth_fixture import auth_db, adapter


def principal(role="advisor", accounts=frozenset({(101, 201)})):
    return AdvisorPrincipal("synthetic-advisor", role, accounts if role == "advisor" else frozenset(), accounts)


def payload(**changes):
    data = json.loads((FIXTURES / "submission.internal.json").read_text(encoding="utf-8"))
    return LandingSubmission.model_validate({**data, **changes})


@pytest.fixture
def env(tmp_path, monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("合成测试禁止网络")

    monkeypatch.setattr(socket, "create_connection", denied)
    monkeypatch.setattr(socket, "getaddrinfo", denied)
    store = SubmissionStore(tmp_path / "synthetic.sqlite")
    store.submit(principal(), payload(), [])
    return store, ShadowService(store)


def create(service, fixture_id="launch"):
    return service.create(principal(), CreateRequest(tenant_id=101, account_id=201,
                          page_id="synthetic-page", submission_revision=1, fixture_id=fixture_id))


def generate(service, **choices):
    state = create(service)
    return service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1,
                        reason="合成顾问确认所有默认及改判来源", choices=choices))


def edit(data, cell, value, sheet="内部规划"):
    book = load_workbook(io.BytesIO(data))
    book[sheet][cell] = value
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def approve_all(service, state):
    return service.approve(principal(), state["plan_id"], ApprovalRequest(expected_version=state["version"],
         row_ids=[r["row_id"] for r in state["rows"]], reason="仅逐项演练复核，未核实项仍保留"))


def test_phase0_connection_and_defaults(env):
    _, service = env
    state = create(service)
    assert state["version"] == 1 and state["state"] == "awaiting_ruling"
    assert not state["p6"]["passed"] and not state["rows"]
    benefit = next(c for c in state["comparisons"] if c["field"] == "benefit")
    assert benefit["conflict"] and len(benefit["candidates"]) == 3
    assert benefit["candidates"][0]["origin"] == "fixture_html"
    assert benefit["default_source_id"] == benefit["candidates"][0]["source_id"]
    assert all(s["artifact_sha256"] and "revision:1" in s["locator"] for s in state["sources"])
    assert state["submission_sha256"] == digest(state["submission"])


@pytest.mark.parametrize("case", ["new_account", "new_product", "event"])
def test_p0_cases_and_p13(env, case):
    store, _ = env
    store.submit(principal(), payload(expected_revision=1, case_type=case), [])
    service = ShadowService(store)
    state = service.create(principal(), CreateRequest(tenant_id=101, account_id=201,
            page_id="synthetic-page", submission_revision=2))
    state = service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1, reason="合成裁定"))
    before = service.expiry(principal(), state["plan_id"], datetime.fromisoformat("2026-11-08T17:59:59+08:00"))
    after = service.expiry(principal(), state["plan_id"], datetime.fromisoformat("2026-11-08T18:00:00+08:00"))
    assert not before[0]["due"] and after[0]["due"]
    assert after[0]["source_refs"] and not after[0]["automatic_pause"]


@pytest.mark.parametrize("changes", [{"case_type": "page_change"}, {"account_ready": False}])
def test_p0_revalidated_at_service_even_if_store_is_corrupted(env, monkeypatch, changes):
    store, service = env
    original = store.read(principal(), 101, 201, "synthetic-page", 1)
    original["payload"].update(changes)
    original["payload_sha256"] = digest(original["payload"])
    monkeypatch.setattr(store, "read", lambda *a: original)
    with pytest.raises(ValidationError):
        create(service)


@pytest.mark.parametrize("backend", ["http", "ocr", "vendor", "real", "", None])
def test_no_real_backend(backend):
    with pytest.raises(ValueError):
        fixture_adapters(backend)


def test_fallback_and_both_unusable(env):
    store, service = env
    state = create(service, "fallback")
    assert all(s["origin"] == "submission" for s in state["sources"])
    store.submit(principal(), payload(expected_revision=1, page_content="合成无字段内容"), [])
    with pytest.raises(ValueError, match="缺项"):
        service.create(principal(), CreateRequest(tenant_id=101, account_id=201,
                       page_id="synthetic-page", submission_revision=2, fixture_id="empty"))


def test_ruling_override_and_full_tree(env):
    _, service = env
    state = create(service)
    c = next(c for c in state["comparisons"] if c["field"] == "benefit")
    chosen = c["candidates"][-1]
    state = service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1,
                        choices={"benefit": chosen["source_id"]}, reason="选择提交内容，合成理由"))
    assert state["selected"]["benefit"] == chosen["source_id"]
    assert chosen["value"] in next(r for r in state["rows"] if r["kind"] == "basic_creative")["text"]
    assert {r["kind"] for r in state["rows"]} == {"project", "campaign", "unit", "keyword",
           "basic_creative", "advanced_creative_placeholder", "negative_keyword", "budget", "region",
           "schedule", "bid_cap", "brand_rule", "timing", "goal"}
    assert all(r["source_refs"] for r in state["rows"])
    assert all(r["verification"] == "pending_verification" for r in state["rows"] if r["kind"] in
               {"campaign", "unit", "keyword", "basic_creative", "advanced_creative_placeholder"})
    assert state["rulings"][0]["reason"] and state["history"][0]["version"] == 1


@pytest.mark.parametrize("choices", [{"made_up": "s-1"}, {"benefit": "missing"}, {"benefit": "s-1"}])
def test_invalid_ruling_atomic(env, choices):
    _, service = env
    state = create(service)
    with pytest.raises(ValueError):
        service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1,
                     choices=choices, reason="合成非法裁定"))
    assert service.read(principal(), state["plan_id"])["version"] == 1


def test_p6_reuse_preserves_unknowns_and_budget_cap(env):
    _, service = env
    state = generate(service)
    assert state["p5"]["passed"]
    rules = {i["rule"] for i in state["p6"]["issues"]}
    assert {"platform_limits", "online_facts", "source_complete"} <= rules
    assert "budget_tree" not in rules and "account_cap" not in rules
    assert not state["p6"]["passed"] and not state["p6"]["execution_export_enabled"]


def test_duplicates_cross_unit_and_existing():
    row = PlanRow(row_id="one", parent_id="u1", kind="keyword", text=" ＡＢＣ ", source_refs=["s-1"],
                  verification="pending_verification")
    second = row.model_copy(update={"row_id": "two", "parent_id": "u2", "text": "abc"})
    report = duplicates([row, second], [{"asset_id": "existing", "kind": "keyword", "text": "AbC"}])
    assert not report["passed"]
    assert report["issues"][0]["matches"] == ["existing"]
    assert report["issues"][1]["matches"] == ["one", "existing"]


def test_advisor_roundtrip_preview_apply_approve_and_invalidation(env):
    _, service = env
    state = generate(service)
    plan_id = state["plan_id"]
    assert not state["advisor_review_complete"]
    state = service.approve(principal(), plan_id, ApprovalRequest(expected_version=2,
                            row_ids=[state["rows"][0]["row_id"]], reason="单行复核"))
    assert not state["advisor_review_complete"] and len(state["approvals"]) == 1
    data = service.export(principal(), plan_id, "advisor")
    preview = service.preview(principal(), plan_id, edit(data, "D8", "合成顾问修改关键词"), "advisor")
    assert preview["changes"] and not preview["applied"]
    assert service.read(principal(), plan_id)["rows"][3]["text"] != "合成顾问修改关键词"
    state = service.apply(principal(), plan_id, ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    assert state["version"] == 3 and not state["approvals"]
    assert state["rows"][3]["verification"] == "pending_verification"
    state = approve_all(service, state)
    assert state["advisor_review_complete"] and not state["p6"]["passed"]
    assert state["customer_confirmation"] == "pending" and not state["execution_export_enabled"]
    state = service.rule(principal(), plan_id, RulingRequest(expected_version=3, reason="合成重裁"))
    assert state["version"] == 4 and not state["approvals"]
    assert len(state["history"][-1]["approvals"]) == len(state["rows"])


def test_customer_preview_never_confirmation_or_apply(env):
    _, service = env
    state = generate(service)
    with pytest.raises(ValueError):
        service.export(principal(), state["plan_id"], "customer")
    state = approve_all(service, state)
    data = service.export(principal(), state["plan_id"], "customer")
    preview = service.preview(principal(), state["plan_id"], edit(data, "D8", "合成客户建议"), "customer")
    assert preview["audience"] == "customer" and preview["rehearsal_only"]
    with pytest.raises(ValueError, match="只预览"):
        service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    current = service.read(principal(), state["plan_id"])
    assert current["version"] == 2 and current["customer_confirmation"] == "pending"


@pytest.mark.parametrize("text,rule", [("合成旧型号滤芯", "exact_duplicate"), ("合成滤芯咨询", "exact_duplicate"),
                                     ("合成免费维修 咨询", "negative_overlap")])
def test_edits_recheck_and_hard_block_approval(env, text, rule):
    _, service = env
    state = generate(service)
    data = edit(service.export(principal(), state["plan_id"], "advisor"), "D8", text)
    preview = service.preview(principal(), state["plan_id"], data, "advisor")
    assert any(i["rule"] == rule for i in preview["p5"]["issues"]+preview["p6"]["issues"])
    state = service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    with pytest.raises(ValueError):
        approve_all(service, state)


def test_budget_edit_rechecked_and_hard_block(env):
    _, service = env
    state = generate(service)
    data = edit(service.export(principal(), state["plan_id"], "advisor"), "E5", 200000)
    preview = service.preview(principal(), state["plan_id"], data, "advisor")
    assert {"budget_tree", "account_cap"} <= {i["rule"] for i in preview["p6"]["issues"]}
    state = service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    with pytest.raises(ValueError):
        approve_all(service, state)


def test_balanced_budget_cannot_silently_override_p1(env):
    _, service = env
    state = generate(service)
    data = service.export(principal(), state["plan_id"], "advisor")
    for cell in ("E5", "E6", "E7"):
        data = edit(data, cell, 13000)
    preview = service.preview(principal(), state["plan_id"], data, "advisor")
    rules = {i["rule"] for i in preview["p6"]["issues"]}
    assert "budget_tree" not in rules and "declared_boundary" in rules
    state = service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    with pytest.raises(ValueError):
        approve_all(service, state)


def test_long_term_no_expiry_and_account_isolation(env):
    store, service = env
    store.submit(principal(), payload(expected_revision=1, case_type="new_product", ends_at=None,
                                     no_deadline_explicit=True, negative_keywords=[]), [])
    state = service.create(principal(), CreateRequest(tenant_id=101, account_id=201,
                    page_id="synthetic-page", submission_revision=2))
    state = service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1, reason="合成长投"))
    assert service.expiry(principal(), state["plan_id"], datetime.fromisoformat("2027-01-01T00:00:00+08:00")) == []
    assert next(r for r in state["rows"] if r["kind"] == "negative_keyword")["text"] == "[]"
    with pytest.raises(ScopeDenied):
        service.read(principal(accounts=frozenset({(102, 202)})), state["plan_id"])


@pytest.mark.parametrize("row_ids", [["missing"], ["r-1", "r-1"]])
def test_approval_ids_must_exist_and_be_individual(env, row_ids):
    _, service = env
    state = generate(service)
    with pytest.raises(ValueError):
        service.approve(principal(), state["plan_id"], ApprovalRequest(expected_version=2,
                        row_ids=row_ids, reason="合成非法复核"))
    assert not service.read(principal(), state["plan_id"])["approvals"]


@pytest.mark.parametrize("cell,value,sheet", [
    ("A5", "forged", "内部规划"), ("B5", "forged", "内部规划"), ("C5", "keyword", "内部规划"),
    ("F8", "已核实", "内部规划"), ("G8", '["forged"]', "内部规划"),
    ("H8", "已演练复核", "内部规划"), ("D8", "=1+1", "内部规划"),
    ("D8", "", "内部规划"), ("E5", 1.5, "内部规划"), ("E5", True, "内部规划"),
    ("E5", 0, "内部规划"), ("E8", 100, "内部规划"),
    ("C3", "篡改来源", "来源"), ("B4", 99, "演练绑定"),
    ("B5", "forged", "演练绑定"), ("B7", "customer", "演练绑定"),
    ("B9", True, "演练绑定"), ("B10", "approved", "演练绑定")])
def test_excel_tampering_and_invalid_values_rejected(env, cell, value, sheet):
    _, service = env
    state = generate(service)
    data = service.export(principal(), state["plan_id"], "advisor")
    with pytest.raises(ValueError):
        service.preview(principal(), state["plan_id"], edit(data, cell, value, sheet), "advisor")
    assert service.read(principal(), state["plan_id"])["version"] == 2


def test_excel_literal_formula_source_not_formula(env):
    store, _ = env
    store.submit(principal(), payload(expected_revision=1, page_content="product: =1+1\nbenefit: 合成\nkeywords: @合成"), [])
    service = ShadowService(store)
    state = service.create(principal(), CreateRequest(tenant_id=101, account_id=201,
                 page_id="synthetic-page", submission_revision=2, fixture_id="fallback"))
    state = service.rule(principal(), state["plan_id"], RulingRequest(expected_version=1, reason="合成来源"))
    data = service.export(principal(), state["plan_id"], "advisor")
    book = load_workbook(io.BytesIO(data), data_only=False)
    assert book["内部规划"]["D5"].value == "=1+1"
    assert book["内部规划"]["D5"].data_type == "s"
    assert parse_workbook(data, state, "advisor")


@pytest.mark.parametrize("data", [b"", b"not a zip", b"a"*(MAX_XLSX_BYTES+1)], ids=["empty", "invalid-zip", "oversize"])
def test_invalid_and_oversize_excel(env, data):
    _, service = env
    state = generate(service)
    with pytest.raises(ValueError):
        service.preview(principal(), state["plan_id"], data, "advisor")


@pytest.mark.parametrize("name,content", [("xl/evil.xml", b'<!DOCTYPE test [<!ENTITY x "evil">]>'),
                                         ("xl/evil.rels", b'<r TargetMode="External"/>'),
                                         ("xl/vbaProject.bin", b"synthetic"), ("../outside", b"synthetic"),
                                         ("xl/big.xml", b"a"*(10*1024*1024+1))],
                         ids=["entity", "external-link", "macro", "traversal", "zip-bomb"])
def test_archive_security(env, name, content):
    _, service = env
    state = generate(service)
    out = io.BytesIO(service.export(principal(), state["plan_id"], "advisor"))
    with zipfile.ZipFile(out, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, content)
    with pytest.raises(ValueError):
        service.preview(principal(), state["plan_id"], out.getvalue(), "advisor")


def test_stale_cross_plan_replay_and_external_mutation(env):
    _, service = env
    state, other = generate(service), generate(service)
    data = service.export(principal(), state["plan_id"], "advisor")
    with pytest.raises(ValueError):
        service.preview(principal(), other["plan_id"], data, "advisor")
    preview = service.preview(principal(), state["plan_id"], data, "advisor")
    with pytest.raises(ValueError):
        service.apply(principal(), other["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2, preview_id=preview["preview_id"]))
    with pytest.raises(RevisionConflict):
        service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=3, preview_id=preview["preview_id"]))
    with pytest.raises(ValueError):
        service.preview(principal(), state["plan_id"], data, "advisor")
    other["rows"][0]["text"] = "外部篡改"
    other["comparisons"].clear()
    fresh = service.read(principal(), other["plan_id"])
    assert fresh["rows"][0]["text"] != "外部篡改" and fresh["comparisons"]


def test_upstream_revision_invalidates_every_old_operation(env):
    store, service = env
    state = generate(service)
    store.submit(principal(), payload(expected_revision=1, goal="合成来源更新"), [])
    with pytest.raises(RevisionConflict):
        service.read(principal(), state["plan_id"])
    with pytest.raises(RevisionConflict):
        approve_all(service, state)
    with pytest.raises(RevisionConflict):
        create(service)


def test_concurrent_revision_only_one_apply(env):
    _, service = env
    state = generate(service)
    data = service.export(principal(), state["plan_id"], "advisor")
    preview = service.preview(principal(), state["plan_id"], data, "advisor")

    def apply():
        try:
            return service.apply(principal(), state["plan_id"], ApplyRequest(expected_version=2,
                                 preview_id=preview["preview_id"]))["version"]
        except RevisionConflict:
            return "conflict"

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(lambda _: apply(), range(2)), key=str) == [3, "conflict"]


def test_audit_hash_chain_and_approval_binding(env):
    _, service = env
    state = approve_all(service, generate(service))
    previous = None
    for event in state["audit"]:
        record = {k: v for k, v in event.items() if k != "sha256"}
        assert event["sha256"] == digest(record) and event["previous_sha256"] == previous
        previous = event["sha256"]
    assert all(a["plan_sha256"] == state["plan_sha256"] and a["version"] == state["version"]
               for a in state["approvals"].values())


@pytest.mark.parametrize("function", [export_execution, enable_plan])
def test_all_execution_functions_unconditionally_closed(function):
    with pytest.raises(PermissionError):
        function(approved=True, force=True, mode="live")


def client_for(service, auth_db, identity=SessionIdentity(1, "advisor", 101, 10)):
    app = FastAPI()
    dependency = principal_dependency(adapter(auth_db), lambda: identity)
    app.include_router(create_router(service, dependency))
    return TestClient(app)


def test_api_complete_rehearsal(env, auth_db):
    _, service = env
    client = client_for(service, auth_db)
    root = "/api/v1/sem/phase1a/shadow/plans"
    response = client.post(root, json={"tenant_id": 101, "account_id": 201,
                            "page_id": "synthetic-page", "submission_revision": 1})
    assert response.status_code == 201
    pid = response.json()["plan_id"]
    response = client.post(f"{root}/{pid}/rulings", json={"expected_version": 1, "reason": "合成裁定"})
    assert response.status_code == 200
    state = response.json()
    data = client.get(f"{root}/{pid}/excel/advisor")
    assert data.status_code == 200 and "not-for-baidu" in data.headers["content-disposition"]
    preview = client.post(f"{root}/{pid}/previews/advisor", files={"file": ("internal.xlsx", data.content)})
    assert preview.status_code == 200
    applied = client.post(f"{root}/{pid}/apply-preview", json={"expected_version": 2,
                           "preview_id": preview.json()["preview_id"]})
    assert applied.status_code == 200
    assert client.post(f"{root}/{pid}/approvals", json={"expected_version": 3,
                       "row_ids": [r["row_id"] for r in state["rows"]], "reason": "合成逐项复核"}).status_code == 200
    customer = client.get(f"{root}/{pid}/excel/customer")
    assert customer.status_code == 200
    response = client.post(f"{root}/{pid}/previews/customer", files={"file": ("sample.xlsx", customer.content)})
    assert response.status_code == 200 and response.json()["customer_confirmation"] == "pending"
    assert client.get(f"{root}/{pid}/expiry-suggestions?now=2026-11-09T00:00:00%2B08:00").json()[0]["due"]
    for path in ("execute", "enable", "baidu-export", "customer-confirm"):
        assert client.post(f"{root}/{pid}/{path}", json={}).status_code == 404


@pytest.mark.parametrize("identity", [SessionIdentity(2, "customer", 101),
    SessionIdentity(1, "advisor", 102, 20), SessionIdentity(3, "customer", 102)])
def test_permissions_for_every_api_operation(env, auth_db, identity):
    _, service = env
    state = generate(service)
    pid = state["plan_id"]
    client = client_for(service, auth_db, identity)
    root = "/api/v1/sem/phase1a/shadow/plans"
    assert client.post(root, json={"tenant_id": 101, "account_id": 201,
                       "page_id": "synthetic-page", "submission_revision": 1}).status_code == 403
    assert client.get(f"{root}/{pid}").status_code == (200 if identity.user_id == 2 else 403)
    assert client.get(f"{root}/{pid}/excel/advisor").status_code == 403
    assert client.get(f"{root}/{pid}/expiry-suggestions").status_code == (200 if identity.user_id == 2 else 403)
    assert client.post(f"{root}/{pid}/rulings", json={"expected_version": 2, "reason": "合成"}).status_code == 403
    assert client.post(f"{root}/{pid}/approvals", json={"expected_version": 2,
                        "row_ids": ["r-1"], "reason": "合成"}).status_code == 403
    assert client.post(f"{root}/{pid}/apply-preview", json={"expected_version": 2,
                        "preview_id": "synthetic"}).status_code == 403
    assert client.post(f"{root}/{pid}/previews/advisor", files={"file": ("sample.xlsx", b"synthetic")}).status_code == 403


def test_api_invalid_inputs_and_stale_conflict(env, auth_db):
    _, service = env
    state = generate(service)
    client = client_for(service, auth_db)
    root = "/api/v1/sem/phase1a/shadow/plans"
    assert client.get(root+"/missing").status_code == 404
    assert client.post(f"{root}/{state['plan_id']}/rulings", json={"expected_version": 1, "reason": "过期"}).status_code == 409
    assert client.get(f"{root}/{state['plan_id']}/excel/baidu").status_code == 422
    assert client.post(root, json={"tenant_id": 101, "account_id": 201, "page_id": "synthetic-page",
                       "submission_revision": 1, "backend": "http"}).status_code == 422
    assert client.post(f"{root}/{state['plan_id']}/previews/advisor",
                       files={"file": ("execution.csv", b"synthetic")}).status_code == 422
    assert client.get(f"{root}/{state['plan_id']}/expiry-suggestions?now=2026-11-09T00:00:00").status_code == 422
