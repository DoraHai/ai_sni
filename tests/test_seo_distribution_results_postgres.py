"""Verified server events -> native publication/evidence; isolated real PostgreSQL."""
import asyncio
import ast
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import json
import os
import subprocess
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import text, select, func
from sqlalchemy.schema import CreateTable
from sqlalchemy.exc import DBAPIError

from app import seo_distribution_results as results
from app import seo_distribution_sources as source
from app.config import get_settings
from app.models.seo import SeoContentAsset, SeoContentPublication, SeoPublishAttempt, SeoDistributionVariant
from app.models.seo_page_capture import SeoPageCapture
from app.models.seo_site_analytics import SeoSiteExportTemplate
from app.api.seo_page_captures import export_publication_list
from test_seo_distribution_sources_postgres import fixture
from test_seo_distribution_media_postgres import images_fixture, freeze
from test_seo_workflow_postgres import ADVISOR

pytestmark=pytest.mark.skipif(not os.getenv('SEO_WORKFLOW_TEST_DATABASE_URL'),reason='isolated local PostgreSQL required')


@pytest.fixture(autouse=True)
def local_capture_only(monkeypatch):
    monkeypatch.setattr(get_settings(),'seo_page_capture_enabled',True)
    async def forbidden(*args,**kwargs): pytest.fail('No external network or page worker during writeback')
    monkeypatch.setattr(httpx.AsyncHTTPTransport,'handle_async_request',forbidden)
    monkeypatch.setattr('app.api.seo_page_captures.execute_page_capture',forbidden)


@asynccontextmanager
async def database(images=False):
    async with (images_fixture() if images else fixture()) as sessions:
        async with sessions.begin() as db:
            connection=await db.connection()
            raw=await connection.get_raw_connection()
            await raw.driver_connection.execute((Path(__file__).parents[1]/'scripts/seo_distribution_results_schema.sql').read_text())
            await db.execute(CreateTable(SeoSiteExportTemplate.__table__,include_foreign_key_constraints=[]))
        yield sessions


async def event(sessions,images=False):
    async with sessions() as db:
        snapshot=await source.load_source(db,7,4,2,60,'xiaohongshu-image' if images else 'wechat-article')
    now=int(datetime.now(timezone.utc).timestamp()*1000)
    return {'event_id':str(uuid4()),'task_id':str(uuid4()),'source_revision':snapshot['source_revision'],
        'content_digest':results.content_digest(snapshot['content']),
        'platform':'xiaohongshu-image' if images else 'wechat-article',
        'public_url':'https://www.xiaohongshu.com/explore/abcdef0123456789' if images else 'https://mp.weixin.qq.com/s/synthetic_article',
        'reported_at':now-1000,'verified_at':now,'verified_by':7,'verification':'human-verified'}


async def record(db,payload,**overrides):
    return await results.record_publication(db,**({'actor_user_id':7,'tenant_id':4,'scope_id':2,'variant_id':60,
        'publication':payload}|overrides))


async def counts(db):
    return tuple([await db.scalar(select(func.count()).select_from(model))
        for model in (SeoContentPublication,SeoPublishAttempt,SeoPageCapture)])


@pytest.mark.parametrize('verification',['human-verified','machine-verified'])
@pytest.mark.parametrize('images',[False,True])
def test_real_native_records_evidence_export_and_replay(verification,images):
    async def scenario():
        async with database(images) as sessions:
            if images: await freeze(sessions)
            payload=await event(sessions,images);payload['verification']=verification
            async with sessions.begin() as db:
                result=await record(db,payload)
                assert result['module']=='seo' and result['publication_id']>0
                assert result['verification']==verification and result['idempotent'] is False
                assert result['monitor_state']=='pending' and result['page_check']['state']=='pending'
                assert result['search_engine_indexed'] is None and result['seo_effect']=='not_evaluated'
                assert await counts(db)==(1,1,1)
                # No internal commit: another connection cannot see any new record.
                async with sessions() as observer: assert await counts(observer)==(0,0,0)
            async with sessions.begin() as db:
                replay=await record(db,payload)
                assert replay['publication_id']==result['publication_id'] and replay['idempotent'] is True
                alias=await record(db,{**payload,'event_id':str(uuid4())})
                assert alias['publication_id']==result['publication_id'] and alias['idempotent'] is True
                assert await counts(db)==(1,1,1)
                publication=await db.get(SeoContentPublication,result['publication_id'])
                attempt=await db.get(SeoPublishAttempt,result['attempt_id'])
                capture=await db.get(SeoPageCapture,result['page_check']['capture_id'])
                assert publication.tenant_id==4 and publication.content_asset_id==40 and publication.variant_id==60
                assert publication.page_url==payload['public_url'] and publication.status=='published'
                assert publication.adapted_content=='<p>selection article body</p>'
                assert publication.link_discovery['verification']==verification
                assert publication.link_discovery['time_basis']=='reported_at'
                assert attempt.response_summary['verified_by']==7 and attempt.response_summary['verification']==verification
                assert capture.relation_type=='publication' and capture.relation_id==publication.id
                assert capture.status=='pending' and capture.source=='auto'
                assert (await db.get(SeoContentAsset,40)).status=='published'
                assert (await db.get(SeoDistributionVariant,60)).status=='approved'
                await source.load_source(db,7,4,2,60,payload['platform'])
                month=datetime.fromtimestamp(payload['reported_at']/1000,timezone.utc).astimezone(
                    __import__('zoneinfo').ZoneInfo('Asia/Shanghai')).strftime('%Y-%m')
                response=await export_publication_list(4,2,month,db,ADVISOR)
                sheet=load_workbook(BytesIO(response.body)).worksheets[0]
                cells=[cell.value for row in sheet for cell in row]
                assert payload['public_url'] in cells and '无截图' in cells
            async with sessions() as db:
                await db.execute(text('SET TRANSACTION READ ONLY'))
                pending=await results.pending_page_checks(db,7,4,2)
                assert pending==[{'capture_id':capture.id,'publication_id':publication.id,'source_url':payload['public_url'],'status':'pending'}]
                assert await counts(db)==(1,1,1)
            async with sessions.begin() as db:
                await db.execute(text("UPDATE seo_page_captures SET status='failed',error_code='timeout' WHERE id=:id"),{'id':capture.id})
            async with sessions() as db:
                assert await results.pending_page_checks(db,7,4,2)==[]
                response=await export_publication_list(4,2,month,db,ADVISOR)
                cells=[cell.value for row in load_workbook(BytesIO(response.body)).worksheets[0] for cell in row]
                assert '失败：截图超时' in cells
    asyncio.run(scenario())


@pytest.mark.parametrize('change',[
    {'tenant_id':9},{'scope_id':99},{'actor_user_id':99},{'variant_id':99},{'variant_id':True},
])
def test_wrong_scope_identity_or_variant_cannot_write(change):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db:
                with pytest.raises(HTTPException): await record(db,payload,**change)
                assert await counts(db)==(0,0,0)
    asyncio.run(scenario())


@pytest.mark.parametrize('change',[
    {'public_url':'https://mp.weixin.qq.com.evil.test/s/a'},
    {'public_url':'https://mp.weixin.qq.com/s/a?ticket=secret'},
    {'public_url':'https://mp.weixin.qq.com:443/s/a'},
    {'public_url':'http://mp.weixin.qq.com/s/a'},
    {'public_url':'https://127.0.0.1/s/a'},
    {'public_url':'https://mp.weixin.qq.com/s/a\n'},
    {'public_url':'https://www.xiaohongshu.com/explore/abcd'},
    {'verified_by':8},{'verified_by':True},{'verification':'unknown'},
    {'verification':'unverified'},{'event_id':'not-uuid'},
    {'content_digest':'0'*64},{'source_revision':'seo-fake'},
    {'reported_at':True},{'verified_at':1},{'verified_at':4102444800000},
    {'body':'client replacement'},{'credential':'secret'},
])
def test_reject_unverified_tampered_or_extra_payload(change):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db:
                with pytest.raises(HTTPException): await record(db,{**payload,**change})
                assert await counts(db)==(0,0,0)
                assert await db.scalar(text('SELECT count(*) FROM seo_distribution_result_events'))==0
    asyncio.run(scenario())


@pytest.mark.parametrize('sql',[
    "UPDATE seo_distribution_variants SET status='draft' WHERE id=60",
    "UPDATE seo_content_assets SET status='review' WHERE id=40",
    "UPDATE seo_content_confirmations SET decision='reject' WHERE id=70",
    "UPDATE seo_site_advisor_assignments SET active=false WHERE id=3",
    "UPDATE users SET is_active=false WHERE id=7",
    "UPDATE roles SET permissions='{}' WHERE id=5",
    "UPDATE seo_distribution_variants SET content='selection changed' WHERE id=60",
])
@pytest.mark.parametrize('already_recorded',[False,True])
def test_current_review_source_and_permissions_required_even_on_retry(sql,already_recorded):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            if already_recorded:
                async with sessions.begin() as db: await record(db,payload)
            async with sessions.begin() as db: await db.execute(text(sql))
            async with sessions.begin() as db:
                with pytest.raises(HTTPException): await record(db,payload)
                assert await counts(db)==((1,1,1) if already_recorded else (0,0,0))
    asyncio.run(scenario())


@pytest.mark.parametrize('aliases',[False,True])
def test_concurrent_events_and_task_aliases_create_one_monitor(aliases):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async def worker():
                async with sessions.begin() as db:
                    return await record(db,{**payload,'event_id':str(uuid4())} if aliases else payload)
            written=await asyncio.gather(*(worker() for _ in range(8)))
            assert len({r['publication_id'] for r in written})==1
            assert sum(not r['idempotent'] for r in written)==1
            async with sessions() as db:
                assert await counts(db)==(1,1,1)
                assert await db.scalar(text('SELECT count(*) FROM seo_distribution_result_events'))==(8 if aliases else 1)
    asyncio.run(scenario())


@pytest.mark.parametrize('new_event',[False,True])
def test_conflicting_duplicate_payload_rejected(new_event):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db: await record(db,payload)
            payload={**payload,'public_url':'https://mp.weixin.qq.com/s/different'}
            if new_event: payload['event_id']=str(uuid4())
            async with sessions.begin() as db:
                with pytest.raises(HTTPException) as exc: await record(db,payload)
                assert exc.value.detail['code']=='seo_distribution_result_'+('task_conflict' if new_event else 'event_conflict')
                assert await counts(db)==(1,1,1)
    asyncio.run(scenario())


def test_caller_rollback_and_atomic_capture_failure(monkeypatch):
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions() as db:
                await record(db,payload)
                await db.rollback()
            async with sessions() as db: assert await counts(db)==(0,0,0)
            async def fail(*args,**kwargs): raise TimeoutError('synthetic secret must not leak')
            monkeypatch.setattr('app.api.seo_page_captures.reserve_publication_page_capture',fail)
            async with sessions.begin() as db:
                with pytest.raises(HTTPException) as exc: await record(db,payload)
                assert exc.value.status_code==503 and 'secret' not in str(exc.value.detail)
                assert await counts(db)==(0,0,0)
                assert (await db.get(SeoContentAsset,40)).status=='ready'
                assert await db.scalar(text('SELECT count(*) FROM seo_distribution_result_events'))==0
                # The caller's outer transaction remains usable after the savepoint rollback.
                assert await db.scalar(text('SELECT 1'))==1
    asyncio.run(scenario())


def test_disabled_capture_is_explicit_and_no_successful_page_claim(monkeypatch):
    monkeypatch.setattr(get_settings(),'seo_page_capture_enabled',False)
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db:
                result=await record(db,payload)
                assert result['monitor_state']=='not_queued'
                assert result['page_check']=={'capture_id':None,'state':'not_queued','reason':'capture_disabled'}
                assert result['search_engine_indexed'] is None and await counts(db)==(1,1,0)
                assert await results.pending_page_checks(db,7,4,2)==[]
    asyncio.run(scenario())


def test_pending_orm_changes_cannot_substitute_server_source():
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions() as db:
                row=await db.get(SeoDistributionVariant,60)
                row.content='client substitution'
                with pytest.raises(HTTPException) as exc: await record(db,payload)
                assert exc.value.detail['code']=='seo_distribution_result_session_not_clean'
                with db.no_autoflush: assert await counts(db)==(0,0,0)
                await db.rollback()
    asyncio.run(scenario())


def test_digest_matches_frozen_central_algorithm():
    central=Path('D:/SNIPERS国内版/ai_sni-worktrees/distribution-control-20261011')
    revision='cb305e8580602569485e55588760cfbd94dc67fc'
    raw=subprocess.check_output(['git','-C',str(central),'show',revision+':app/distribution_contract.py'])
    node=next(n for n in ast.parse(raw).body if isinstance(n,ast.FunctionDef) and n.name=='content_digest')
    namespace={'hashlib':hashlib,'json':json}
    exec(compile(ast.Module(body=[node],type_ignores=[]),'frozen-central-digest','exec'),namespace)
    for assets in ([],[{'id':'asset-一','sha256':'a'*64,'mediaType':'image/png','bytes':1234}]):
        content={'version':'seo-m3-v2-hash','title':'中文 🐯','body':'selection\n\\"\u2028','assets':assets}
        assert results.content_digest(content)==namespace['content_digest'](content)


def test_different_tasks_cannot_adopt_same_native_publication():
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async def worker(p):
                async with sessions.begin() as db:
                    try: return await record(db,p)
                    except HTTPException as exc: return exc
            returned=await asyncio.gather(worker(payload),worker({**payload,'event_id':str(uuid4()),'task_id':str(uuid4())}))
            assert sum(isinstance(r,dict) for r in returned)==1
            assert [r.detail['code'] for r in returned if isinstance(r,HTTPException)]==['seo_distribution_result_publication_conflict']
            async with sessions() as db: assert await counts(db)==(1,1,1)
    asyncio.run(scenario())


def test_ledger_is_immutable_and_missing_schema_fails_closed():
    async def scenario():
        async with database() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db: await record(db,payload)
            for sql in ("UPDATE seo_distribution_result_events SET result='{}'",'DELETE FROM seo_distribution_result_events'):
                async with sessions.begin() as db:
                    with pytest.raises(DBAPIError):
                        async with db.begin_nested(): await db.execute(text(sql))
                    assert await db.scalar(text('SELECT count(*) FROM seo_distribution_result_events'))==1
        async with fixture() as sessions:
            payload=await event(sessions)
            async with sessions.begin() as db:
                with pytest.raises(HTTPException) as exc: await record(db,payload)
                assert exc.value.status_code==503 and await counts(db)==(0,0,0)
    asyncio.run(scenario())
