"""Real HTTP/SQL projections over two synthetic tenants; all external I/O mocked."""
import asyncio
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import load_workbook

from test_geo_project_client import env, get, add_answers, NOW, WEEK, Answer, Run, Publication, Unit
from app.geo.content.engine_identity import trusted_engine_source
from app.models import GeoActionTicket


def test_open_week_progress_is_project_scoped_real_only_and_separate_from_closed_week(env):
    from test_geo_project_client import Prompt
    from app.geo.content import time_windows
    now = datetime(2026, 10, 22, 4)
    start = datetime(2026, 10, 18, 16)
    env.monkeypatch.setattr(time_windows, 'now_utc_naive', lambda: now)
    add_answers(env)
    probe = env.session.get(Prompt, 3)
    probe.is_brand_probe = True
    rows = []
    cases = [
        (1, 101, start, 'openai_compat', False, ''),
        (1, 101, now-timedelta(minutes=1), 'openai_compat', False, 'analysis=needs_review'),
        (2, 101, now-timedelta(minutes=1), 'openai_compat', False, ''),
        (2, 101, now-timedelta(minutes=1), 'unknown', False, '真采样'),
        (3, 101, start, 'openai_compat', False, ''),  # brand probe
        (1, 202, start, 'openai_compat', False, ''),  # mismatched tenant on authorized prompt
        (4, 101, start, 'openai_compat', False, ''),  # another project
        (5, 202, start, 'openai_compat', False, ''),
        (6, 101, start, 'openai_compat', False, ''),  # unbound prompt
        (7, 202, start, 'openai_compat', False, ''),  # foreign tenant on bound unit
        (1, 101, start, 'openai_compat', True, ''),
        (1, 101, start, 'mock_persona', False, ''),
        (1, 101, start, 'manual', False, 'openai_compat'),
        (1, 101, start, 'unknown', False, ''),
        (1, 101, start-timedelta(microseconds=1), 'openai_compat', False, ''),
        (1, 101, now, 'openai_compat', False, ''),
        (1, 101, now+timedelta(hours=1), 'openai_compat', False, ''),
        (1, 101, None, 'openai_compat', False, ''),
    ]
    for i, (pid, tid, captured, mode, simulated, note) in enumerate(cases, 201):
        rows.append(Answer(id=i, tenant_id=tid, prompt_id=pid, engine='doubao',
            captured_at=captured, sample_mode=mode, simulated=simulated, note=note, raw_text='Synthetic'))
    env.save(probe, *rows)
    expected = {'week_start': '2026-10-19', 'next_count_date': '2026-10-26',
                'detected_answer_count': 4, 'basis_code': 'saved_real_api_answers_open_week'}
    for suffix in ('/overview', '/visibility', '/summary', '/questions', '/questions/1'):
        for week in (WEEK, '2026-10-12'):
            response = get(env, suffix, week_end=week)
            assert response.status_code == 200, response.text
            assert response.json()['data']['open_week_progress'] == expected
    questions = get(env, '/questions').json()['data']['items']
    assert {q['id']: q['open_week_detected_count'] for q in questions} == {1: 2, 2: 2, 3: 0}
    assert get(env, '/questions/1').json()['data']['question']['open_week_detected_count'] == 2
    assert get(env, '/summary', project=2, week_end=WEEK).json()['data']['open_week_progress']['detected_answer_count'] == 1
    assert get(env, '/summary', project=3, tenant=202, week_end=WEEK).status_code == 403
    env.ctx.is_superadmin = True
    env.ctx.tenant_id = None
    assert get(env, '/summary', project=3, tenant=202, week_end=WEEK).json()['data']['open_week_progress']['detected_answer_count'] == 1
    # No current-week rows enter historical basis, counts or matrix observations.
    result = get(env, '/overview').json()
    assert result['basis']['source_counts']['real'] == 9  # eight fixtures plus just-before-Monday boundary
    assert result['data']['monitoring']['saved_answer_count'] == 9
    assert get(env, '/summary', week_end=WEEK).json()['data']['qualified_counts']['samples'] == 6
    assert {q['id']: q['observation']['total'] for q in questions} == {1: 3, 2: 3, 3: 2}


@pytest.mark.parametrize('project', [1, 4])
def test_open_week_zero_always_present_for_empty_or_unbound_project(env, project):
    for suffix in ('/overview', '/visibility', '/summary', '/questions'):
        result = get(env, suffix, project=project, week_end=WEEK)
        assert result.status_code == 200
        assert result.json()['data']['open_week_progress']['detected_answer_count'] == 0
    if project == 1:
        data = get(env, '/questions/1').json()['data']
        assert data['open_week_progress']['detected_answer_count'] == data['question']['open_week_detected_count'] == 0


@pytest.mark.parametrize('now,week', [
    (datetime(2026, 10, 18, 15, 59, 59), date(2026, 10, 12)),
    (datetime(2026, 10, 18, 16), date(2026, 10, 19)),
    (datetime(2027, 1, 1, tzinfo=timezone.utc), date(2026, 12, 28)),
])
def test_open_week_uses_shanghai_monday_across_utc_and_year_boundaries(now, week):
    from app.geo.content.time_windows import open_week_window, to_utc_naive
    actual, start, stop = open_week_window(now)
    assert actual == week and stop == to_utc_naive(now)
    assert start == datetime.combine(week, datetime.min.time()) - timedelta(hours=8)


def test_open_week_query_is_bounded_and_never_falls_back_to_customer(env):
    from app.geo import project_client_store as store
    env.monkeypatch.setattr(store, 'ROW_BUDGET', 4)
    env.monkeypatch.setattr(store, 'open_week_window', lambda: (date(2026, 10, 19),
        datetime(2026, 10, 18, 16), datetime(2026, 10, 22, 4)))
    env.save(*[Answer(id=201+i, tenant_id=101, prompt_id=1, engine='doubao',
        captured_at=datetime(2026, 10, 20), sample_mode='openai_compat', simulated=False, raw_text='Synthetic')
        for i in range(5)])
    result = get(env, '/summary', week_end=WEEK)
    assert result.status_code == 503
    assert result.json()['detail']['code'] == 'GEO_CLIENT_READ_UNAVAILABLE'
    assert get(env, '/summary', project=2, week_end=WEEK).json()['data']['open_week_progress']['detected_answer_count'] == 0


@pytest.mark.parametrize('reason', ['engine_source_untrusted', 'analysis_incomplete'])
def test_matrix_tested_excluded_reports_real_reason_and_capture_time(env, reason):
    row = add_answers(env, 1)[0]
    run = env.session.get(Run, 99)
    cells = deepcopy(run.items)
    if reason == 'engine_source_untrusted':
        cells[0].update(provider='deepseek', model='deepseek-chat')
    else:
        row.note = 'method=unprimed_json_v2 analysis=needs_review'
        cells[0]['analysis_status'] = 'needs_review'
    run.items = cells
    env.save(row, run)
    for suffix in ('/questions', '/questions/1'):
        data = get(env, suffix).json()['data']
        question = next(q for q in data['items'] if q['id'] == 1) if suffix == '/questions' else data['question']
        cell = question['by_engine'][0]
        assert cell['status'] == 'tested_excluded'
        assert reason in cell['excluded_reason_codes']
        assert len(cell['excluded_reason_codes']) == len(set(cell['excluded_reason_codes']))
        assert cell['last_sample_at'] == '2026-10-14T05:00:00Z'
        assert cell['percentage'] is None and cell['total'] is None and cell['mentioned'] is None
        assert all(c['excluded_reason_codes'] == [] for c in question['by_engine'][1:])
    workbook = load_workbook(BytesIO(get(env, '/reports/export.xlsx', week_end=WEEK).content))
    assert workbook['问题矩阵']['I2'].value == '已检测，未计入'
    # Inspect the actual scoped PDF HTML before replacing browser rendering.
    html = []
    async def render(content):
        html.append(content)
        return b'%PDF-synthetic'
    env.monkeypatch.setattr('app.geo.report_pdf.render_report_pdf', render)
    assert get(env, '/reports/export.pdf', week_end=WEEK).status_code == 200
    assert '已检测，未计入' in html[0]


def test_matrix_unavailable_snapshot_is_scoped_and_latest_attempt_wins(env):
    env.save(Run(id=80, tenant_id=101, status='completed', started_at=NOW, finished_at=NOW,
        items=[{'prompt_id': 1, 'engine': 'doubao', 'ok': True, 'snapshot_id': 9999,
                'raw_text': 'FOREIGN_SNAPSHOT_TEXT'},
               {'prompt_id': 4, 'engine': 'qwen', 'ok': True, 'snapshot_id': 9998}]))
    result = get(env, '/questions/1')
    cell = result.json()['data']['question']['by_engine'][0]
    assert cell['status'] == 'tested_excluded' and cell['excluded_reason_codes'] == ['snapshot_unavailable']
    assert cell['last_sample_at'] is None and 'FOREIGN_SNAPSHOT_TEXT' not in result.text
    assert get(env, '/questions/1', week_end='2026-10-12').json()['data']['question']['by_engine'][0]['status'] == 'not_tested'
    env.save(Run(id=81, tenant_id=101, status='completed', started_at=NOW+timedelta(hours=1),
        finished_at=NOW+timedelta(hours=1), items=[{'prompt_id': 1, 'engine': 'doubao', 'ok': False}]))
    cell = get(env, '/questions/1').json()['data']['question']['by_engine'][0]
    assert cell['status'] == 'failed' and cell['excluded_reason_codes'] == []
    env.save(Run(id=82, tenant_id=101, status='completed', started_at=NOW+timedelta(hours=1),
        finished_at=NOW+timedelta(hours=1, microseconds=1), items=[
            {'prompt_id': 1, 'engine': 'doubao', 'ok': True, 'snapshot_id': 9999}]))
    cell = get(env, '/questions/1').json()['data']['question']['by_engine'][0]
    assert cell['status'] == 'tested_excluded' and cell['excluded_reason_codes'] == ['snapshot_unavailable']


def test_qualified_matrix_keeps_observed_and_small_sample_despite_latest_excluded_attempt(env):
    rows = add_answers(env, count=8, engines=1, questions=1)
    run = env.session.get(Run, 99)
    cells = deepcopy(run.items)
    cells[-1].update(provider='deepseek', model='deepseek-chat')
    run.items = cells
    env.save(run)
    cell = get(env, '/questions/1').json()['data']['question']['by_engine'][0]
    assert cell['status'] == 'observed' and cell['total'] == 7 and cell['excluded_reason_codes'] == []
    for row in rows[:5]:
        row.note = 'method=unprimed_json_v2 analysis=needs_review'
    env.save(*rows[:5])
    cell = get(env, '/questions/1').json()['data']['question']['by_engine'][0]
    assert cell['status'] == 'small_sample' and cell['total'] == 2 and cell['excluded_reason_codes'] == []


@pytest.mark.parametrize('start,end,label', [
    ('2026-10-12T00:00:00+08:00', '2026-10-19T00:00:00+08:00', '2026-10-12 至 2026-10-18'),
    ('2026-12-27T16:00:00Z', '2027-01-03T16:00:00Z', '2026-12-28 至 2027-01-03'),
    ('2026-09-27T16:00:00Z', '2026-10-04T16:00:00Z', '2026-09-28 至 2026-10-04'),
])
def test_project_report_period_is_inclusive_shanghai_dates(start, end, label):
    from app.geo.report_export import project_period_label
    assert project_period_label({'start': start, 'end': end}) == label


def test_project_report_html_and_xlsx_use_friendly_closed_week_dates(env):
    html = []
    async def render(content):
        html.append(content)
        return b'%PDF-synthetic'
    env.monkeypatch.setattr('app.geo.report_pdf.render_report_pdf', render)
    assert get(env, '/reports/export.pdf', week_end=WEEK).status_code == 200
    label = '统计周期：2026-10-12 至 2026-10-18（北京时间）'
    assert label in html[0] and '右端不含' not in html[0]
    workbook = load_workbook(BytesIO(get(env, '/reports/export.xlsx', week_end=WEEK).content))
    assert label in [r[0] for r in workbook['口径与缺数'].values]


@pytest.mark.parametrize('mutation', ['missing', 'failed', 'wrong_article', 'stale', 'edited_variant'])
def test_customer_approval_cannot_expose_internally_unchecked_content(env, mutation):
    from test_geo_project_client import Task, Variant
    task = env.session.get(Task, 1)
    assert task.review_status == 'approved'
    checks = deepcopy(task.rule_result)
    if mutation == 'missing': checks = None
    if mutation == 'failed': checks['ready'] = False
    if mutation == 'wrong_article': checks['article_id'] = 4
    if mutation == 'stale': checks['checked_at'] = (NOW - timedelta(days=2)).isoformat()
    if mutation == 'edited_variant':
        variant = env.session.get(Variant, 1)
        variant.adapt_meta = {**variant.adapt_meta, 'publishable': False}
        env.save(variant)
    else:
        task.rule_result = checks
        env.save(task)
        assert get(env, '/content/1').status_code == 404
        assert all(r['id'] != 1 for r in get(env, '/content').json()['data']['items'])
    assert get(env, '/content/1/variants/1').status_code == 404
    assert all(r['id'] != 1 for r in get(env, '/article-citations', week_end=WEEK).json()['data']['items'])


def test_search_inventory_and_keyword_suggestions_stay_in_project(env):
    from app.geo.content import search_visibility as sv, search_visibility_routes as routes
    from test_geo_project_client import SelectOnly
    from app.models import GeoProject, GeoMediaPlacement
    from sqlalchemy import event
    # Create this pre-existing model only in the in-memory fixture database.
    event.remove(env.engine, 'before_cursor_execute', env.guard)
    try:
        GeoMediaPlacement.__table__.create(env.engine)
    finally:
        event.listen(env.engine, 'before_cursor_execute', env.guard)
    env.save(*[GeoMediaPlacement(id=i, tenant_id=tid, related_prompt_id=prompt,
        published_url=f'https://media.example.com/{i}', name='合成投放')
        for i, tid, prompt in [(1, 101, 1), (2, 101, 4), (3, 202, 5), (4, 101, None)]])
    session = SelectOnly(env.session)
    for pid, expected in [(1, {'https://p1.example.com/post/1', 'https://p1.example.com/post/2',
                              'https://p1.example.com/post/3', 'https://media.example.com/1'}),
                          (2, {'https://p2.example.com/post/4', 'https://media.example.com/2'}), (4, set())]:
        project = env.session.get(GeoProject, pid)
        assert set(asyncio.run(sv.article_urls(session, project, sv.VisibilityConfig()))) == expected
    unit = env.session.get(Unit, 11)
    other = env.session.get(Unit, 22)
    unit.keyword, other.keyword = '本项目关键词', '另一项目关键词'
    env.save(unit, other)
    env.ctx.role_name = 'advisor'
    env.ctx.permissions = {'geo.assets': 'edit', 'geo.content': 'edit'}
    env.monkeypatch.setattr(routes, 'load_project', AsyncMock(return_value=(
        env.session.get(GeoProject, 1), sv.VisibilityConfig())))
    env.monkeypatch.setattr(routes, 'access_state', lambda *args: {})
    result = asyncio.run(routes.get_config(1, 101, env.ctx, session))
    assert result['suggested_core_terms'] == ['本项目关键词']
    assert result['suggested_core_terms_scope'] == 'project'


def test_failed_and_unconfigured_cells_are_scoped_and_missing_is_null(env):
    env.save(Run(id=80, tenant_id=101, status='completed', started_at=NOW, finished_at=NOW,
        items=[{'prompt_id': 1, 'engine': 'doubao', 'ok': False,
                'skipped_reason': 'skipped:prefer_real_no_platform_key'},
               {'prompt_id': 1, 'engine': 'deepseek', 'ok': False, 'error': 'synthetic failure'},
               {'prompt_id': 4, 'engine': 'qwen', 'ok': False, 'error': 'FOREIGN_FAILURE'}]))
    result = get(env, '/questions/1')
    cells = result.json()['data']['question']['by_engine']
    assert [c['status'] for c in cells[:3]] == ['not_configured', 'failed', 'not_tested']
    assert all(c['percentage'] is None and c['total'] is None for c in cells)
    assert 'FOREIGN_FAILURE' not in result.text


def test_export_escapes_formulas_html_and_unreviewed_recommendations(env):
    from app.onsite_workflow import proposal_hash
    items = [{'kind': 'faq', 'instruction': 'UNREVIEWED_PRIVATE'}]
    env.save(GeoActionTicket(id=60, tenant_id=101, title='UNREVIEWED_PRIVATE',
        advice_code='onsite:v1:unreviewed', status='todo',
        progress={'onsite': {'project_id': 1, 'phase': 'review',
                            'items': items, 'approval': {'hash': proposal_hash(items)}}}))
    from test_geo_project_client import Variant
    variant = env.session.get(Variant, 1)
    variant.title = '=HYPERLINK("synthetic")'
    env.save(variant)
    result = get(env, '/reports/export.xlsx', week_end=WEEK)
    book = load_workbook(BytesIO(result.content))
    assert book['文章引用']['D2'].value.startswith("'=")
    assert 'UNREVIEWED_PRIVATE' not in str([list(ws.values) for ws in book])
    renderer = AsyncMock(return_value=b'%PDF-synthetic')
    env.monkeypatch.setattr('app.geo.report_pdf.render_report_pdf', renderer)
    assert get(env, '/reports/export.pdf', week_end=WEEK).status_code == 200
    assert 'UNREVIEWED_PRIVATE' not in renderer.call_args.args[0]
    assert '元宝数据取自腾讯混元模型接口' in renderer.call_args.args[0]


@pytest.mark.parametrize('engine,provider,model,valid', [
    ('doubao', 'deepseek', 'deepseek-chat', False), ('doubao', 'doubao', 'deepseek-chat', False),
    ('doubao', 'doubao', 'ep-synthetic', True), ('qwen', 'qwen', 'qwen-plus', True),
    ('qwen', 'dashscope', 'deepseek-v3', False), ('deepseek', 'dashscope:deepseek', 'deepseek-v3', True),
    ('deepseek', 'dashscope', 'deepseek-v3', True), ('hunyuan', 'hunyuan', 'hunyuan-turbos', True),
    ('hunyuan', 'hunyuan', 'hy4-preview', True), ('hunyuan', 'hunyuan', 'HY4-Preview', True),
    ('hunyuan', 'hunyuan', 'hy4', False), ('hunyuan', 'hunyuan', 'hy3-preview', False),
    ('deepseek', 'deepseek', 'hy4-preview', False), ('qwen', 'qwen', 'hy4-preview', False),
    ('hunyuan', 'openai', 'hy4-preview', False),
    ('doubao', None, None, False), ('other', 'other', 'model', False)])
def test_identity(engine, provider, model, valid):
    assert trusted_engine_source(engine, provider, model) is valid


@pytest.mark.parametrize('engine,stance,configured,expected_calls', [
    ('doubao', 'hybrid', False, 0), ('doubao', 'real_only', False, 0),
    ('qwen', 'hybrid', True, 1), ('deepseek', 'real_only', True, 1)])
def test_patrol_never_promotes_shared_deepseek(monkeypatch, engine, stance, configured, expected_calls):
    from app.geo.content import patrol, engine_providers, ai_settings
    run = NS(id=99, tenant_id=101, status='pending', trigger='manual', auto_persist=False,
             prefer_real=True, prompt_limit=5, engine_keys=[engine], summary=None, items=None,
             error=None, started_at=None, finished_at=None, created_by=None)
    prompt = NS(id=1, question='合成问题', tags=[], priority=1)
    tenant = NS(id=101, name='合成品牌', brand_terms=['合成品牌'])
    session = AsyncMock()
    session.info = {}
    session.get = AsyncMock(side_effect=lambda model, pk: tenant if model.__name__ == 'Tenant' else run)
    session.scalars = AsyncMock(side_effect=[[NS(id=1, engine_key=engine, enabled=True)], [prompt]])
    session.add = MagicMock()
    llm = {'api_key': 'synthetic', 'base_url': 'https://synthetic.invalid/v1',
           'provider': 'deepseek', 'model': 'deepseek-chat'}
    dedicated = dict(llm, provider=engine if engine == 'qwen' else 'dashscope:deepseek', model=engine+'-test')
    monkeypatch.setattr(patrol, 'ensure_geo_entitlement', AsyncMock())
    monkeypatch.setattr(ai_settings, 'resolve_llm_credentials', AsyncMock(return_value=llm))
    monkeypatch.setattr(ai_settings, 'ensure_ai_setting', AsyncMock(return_value=NS(monitoring_stance=stance)))
    monkeypatch.setattr(engine_providers, 'resolve_platform_engine_credentials', lambda key: dedicated if configured else None)
    probe = AsyncMock(return_value={'raw_text': '合成回答', 'sample_mode': 'openai_compat',
        'simulated': False, 'provider': dedicated['provider'], 'model': dedicated['model'], 'suggested_mentions_brand': True})
    monkeypatch.setattr(patrol, 'run_probe_draft', probe)
    result = asyncio.run(patrol.execute_patrol_run(session, 99, tenant_id=101))
    assert result.status == 'completed'
    assert probe.await_count == expected_calls
    if not configured:
        assert result.items[0]['skipped_reason'].startswith('skipped:')
        assert result.summary['real_samples'] == 0
        session.add.assert_not_called()


def test_summary_uses_one_cohort_positions_and_mentioned_sentiment(env):
    rows = add_answers(env)
    for i, row in enumerate(rows):
        row.brand_position = ('first', 'alternative', 'mentioned', 'absent')[i % 4]
        row.sentiment = ('positive', 'neutral', 'negative', 'unknown')[i % 4]
    env.save(*rows)
    result = get(env, '/summary', week_end=WEEK).json()
    assert result['data']['qualified_counts'] == {'samples': 8, 'questions': 3, 'engines': 2}
    metrics = result['data']['metrics']
    assert metrics['brand_mention_rate'] == 100 and metrics['brand_mention_answer_count'] == 8
    assert metrics['average_recommendation_position'] == 2 and metrics['valid_position_count'] == 6
    assert metrics['brand_positive_rate'] == 33.3333 and metrics['annotated_mention_count'] == 6
    assert [r['engine_key'] for r in result['data']['by_engine']] == ['doubao','deepseek','qwen','hunyuan','wenxin','kimi']
    assert result['data']['by_engine'][3]['label'] == '元宝'
    assert '元宝数据取自腾讯混元模型接口' in result['data']['footnotes']
    # A non-mention labelled positive must not increase the sentiment denominator.
    rows[0].mentions_brand = False
    run = env.session.get(Run, 99)
    cells = deepcopy(run.items)
    cells[0]['suggested_mentions_brand'] = False
    run.items = cells
    env.save(rows[0], run)
    m = get(env, '/summary', week_end=WEEK).json()['data']['metrics']
    assert m['mention_count'] == 7 and m['annotated_mention_count'] == 5 and m['brand_positive_rate'] == 20


def test_untrusted_history_excluded_from_old_and_new_metrics_without_rewrite(env):
    add_answers(env)
    run = env.session.get(Run, 99)
    cells = deepcopy(run.items)
    cells[0].update(provider='deepseek', model='deepseek-chat')  # snapshot labelled doubao
    run.items = cells
    env.save(run)
    for suffix in ('/summary', '/visibility'):
        data = get(env, suffix, week_end=WEEK).json()
        assert {'code': 'engine_source_untrusted', 'count': 1} in data['basis']['exclusion_counts']
        sample = data['data'].get('current', data['data'])
        assert sample['qualified_counts']['samples'] == 7
    assert env.session.get(Answer, 1).engine == 'doubao'


@pytest.mark.parametrize('suffix', ['/summary', '/article-citations', '/recommendations', '/reports/export.xlsx', '/reports/export.pdf'])
def test_new_endpoints_reauthorize_and_cannot_cross_projects(env, suffix):
    # PDF rendering is replaced only after the actual scoped payload is built.
    env.monkeypatch.setattr('app.geo.report_pdf.render_report_pdf', AsyncMock(return_value=b'%PDF-synthetic'))
    assert get(env, suffix, tenant=202, project=3, **({'week_end': WEEK} if suffix != '/recommendations' else {})).status_code == 403
    assert get(env, suffix, project=3).status_code == 404
    assert get(env, suffix, **({'week_end': WEEK} if suffix != '/recommendations' else {})).status_code == 200
    env.ctx.permissions.pop('geo.assets')
    assert get(env, suffix).status_code == 403


def test_matrix_filters_are_server_side_and_cursor_binds_filter(env):
    add_answers(env, 1)
    unit = env.session.get(Unit, 11)
    unit.keyword = '真正关键词'
    env.save(unit)
    data = get(env, '/questions', business_id=11, unit_id=11, limit=1).json()['data']
    assert data['pagination']['total_count'] == 3 and data['items'][0]['unit']['display_name'] == '真正关键词'
    cursor = data['pagination']['next_cursor']
    assert get(env, '/questions', business_id=22).status_code == 404
    assert get(env, '/questions', business_id=11, limit=1, cursor=cursor).status_code == 400
    first = get(env, '/questions/1').json()['data']['question']['by_engine'][0]
    assert first['mentioned'] == 1 and first['total'] == 1 and first['small_sample']
    assert get(env, '/questions/1').json()['data']['question']['by_engine'][1]['percentage'] is None


def test_empty_binding_and_no_samples_are_null(env):
    for project in (1, 4):
        summary = get(env, '/summary', project=project, week_end=WEEK).json()['data']
        assert summary['metrics']['brand_mention_rate'] is None
        assert summary['metrics']['brand_mention_answer_count'] is None
        articles = get(env, '/article-citations', project=project, week_end=WEEK).json()['data']
        assert articles['coverage']['percentage'] is None
    assert get(env, '/questions', project=4).json()['data']['items'] == []


def test_article_exact_query_dedup_publication_time_and_evidence(env):
    rows = add_answers(env)
    publication = env.session.get(Publication, 1)
    publication.published_url = 'https://p1.example.com/post?id=1'
    publication.published_at = NOW - timedelta(hours=1)
    rows[0].raw_text = '回答 https://p1.example.com/post?id=1&utm_source=x'
    rows[0].cited_urls = ['https://p1.example.com/post?id=1&utm_source=x']
    rows[1].raw_text = '回答 https://p1.example.com/post?id=2'
    rows[1].cited_urls = ['https://p1.example.com/post?id=2']
    run = env.session.get(Run, 99)
    cells = deepcopy(run.items)
    for i in (0, 1): cells[i]['raw_text'] = rows[i].raw_text
    run.items = cells
    duplicate = Publication(id=20, variant_id=1, channel='website', status='published',
        published_url='https://p1.example.com/post?id=1&utm_source=duplicate', published_at=NOW)
    env.save(publication, duplicate, rows[0], rows[1], run)
    data = get(env, '/article-citations', week_end=WEEK).json()['data']
    item = next(a for a in data['items'] if a['id'] == 1)
    assert item['mentioned'] == 1 and item['total'] == 8 and item['percentage'] == 12.5
    assert [e['answer_id'] for e in item['evidence']] == [1]
    assert data['coverage']['published_articles'] == 3
    publication.published_at = NOW + timedelta(hours=1)
    duplicate.published_at = NOW + timedelta(hours=1)
    env.save(publication, duplicate)
    item = next(a for a in get(env, '/article-citations', week_end=WEEK).json()['data']['items'] if a['id'] == 1)
    assert item['total'] is None and item['percentage'] is None and item['evidence'] == []


def test_recommendations_allowlist_and_report_contains_same_summary(env):
    from app.onsite_workflow import proposal_hash
    add_answers(env)
    items = [{'kind': 'faq', 'instruction': '补充FAQ', 'private': 'HIDDEN'}]
    env.save(GeoActionTicket(id=51, tenant_id=101, title='本项目建议', advice_code='onsite:v1:synthetic', status='todo',
        progress={'onsite': {'project_id': 1, 'phase': 'implementation', 'approval': {'hash': proposal_hash(items)}, 'items': items}}),
        GeoActionTicket(id=52, tenant_id=101, title='FOREIGN_PROJECT', advice_code='onsite:v1:synthetic2', status='todo',
        progress={'onsite': {'project_id': 2, 'phase': 'draft', 'items': []}}))
    response = get(env, '/recommendations')
    assert response.status_code == 200 and 'HIDDEN' not in response.text and 'FOREIGN_PROJECT' not in response.text
    assert response.json()['data']['items'][0]['recommendations'][0]['description'] == '补充FAQ'
    report = get(env, '/reports/export.xlsx', week_end=WEEK)
    assert report.status_code == 200 and report.headers['cache-control'] == 'private, no-store'
    workbook = load_workbook(BytesIO(report.content))
    assert workbook['四项汇总']['C2'].value == get(env, '/summary', week_end=WEEK).json()['data']['metrics']['brand_mention_rate']
    assert workbook['问题矩阵'].max_row == 19
    assert '元宝数据取自腾讯混元模型接口' in [r[0] for r in workbook['口径与缺数'].values]
    assert 'FOREIGN_PROJECT' not in str([list(ws.values) for ws in workbook])
