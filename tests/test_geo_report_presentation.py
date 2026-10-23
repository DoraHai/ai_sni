"""Pure export regressions: synthetic snapshot only, no database or settings."""
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.geo.report_export import (
    build_project_report_html, build_project_xlsx, report_beijing_time,
)


def snapshot():
    stamp = '2026-10-15T16:05:00Z'
    metrics = dict(brand_mention_rate=50, brand_mention_answer_count=4,
                   average_recommendation_position=1.5, brand_positive_rate=50,
                   sample_count=8, mention_count=4, valid_position_count=4,
                   position_distribution=dict(first=2, alternative=2, mentioned=0),
                   annotated_mention_count=4)
    return dict(project=dict(id=1, name='<项目>'),
                summary=dict(metrics=metrics, by_engine=[], sample_status='ready', footnotes=[],
                             period=dict(start='2026-10-12T00:00:00+08:00', end='2026-10-19T00:00:00+08:00')),
                questions=[dict(id=1, business=dict(name='业务'), unit=dict(display_name='关键词'), question='问题',
                                by_engine=[dict(label='元宝', mentioned=3, total=3, percentage=100,
                                                status='observed', last_sample_at=stamp)])],
                citations=dict(coverage=dict(cited_articles=1, published_articles=2, percentage=50, status='observed'),
                               items=[dict(id=1, task_id=2, article_id=3, title='文章', url='https://example.com/a',
                                           published_at=stamp, mentioned=1, total=2, percentage=50, status='observed',
                                           evidence=[dict(answer_id=4, question_id=1, engine_key='hunyuan',
                                                          captured_at=stamp, answer_path='/answers/4')])]),
                recommendations=[dict(id=5, title='建议', phase='doing', next_action='检查', checked_at=stamp,
                                      recommendations=[dict(title='补全', description='说明', priority='high', status='open')])],
                basis=dict(last_data_at=stamp, exclusion_counts=[]))


def test_pdf_and_excel_share_beijing_timestamps_and_chinese_labels():
    data = snapshot()
    html = build_project_report_html(data)
    workbook = load_workbook(BytesIO(build_project_xlsx(data)))
    stamp = '2026-10-16 00:05:00 北京时间'
    assert workbook['问题矩阵']['J2'].value == stamp
    assert workbook['文章引用']['F2'].value == stamp
    assert workbook['引用依据']['E2'].value == stamp
    assert workbook['站内建议']['E2'].value == stamp
    assert workbook['引用依据']['D2'].value == '元宝'
    assert [workbook['站内建议'].cell(2, c).value for c in (3, 8, 9)] == ['进行中', '高', '待处理']
    assert html.count(stamp) == 5
    for label in ('进行中', '高', '待处理', '元宝', '2026-10-12 至 2026-10-18（北京时间）'):
        assert label in html
    for raw in ('2026-10-15T16:05:00Z', '>doing<', '>high<', '>open<', '>hunyuan<'):
        assert raw not in html
    assert '&lt;项目&gt;' in html
    assert workbook['文章引用']['I2'].value == 50


def test_missing_times_remain_unknown_and_offsets_and_naive_utc_are_supported():
    assert report_beijing_time(None) == '暂无'
    assert report_beijing_time('invalid') == '暂无'
    assert report_beijing_time('2026-10-16T00:05:00+08:00') == '2026-10-16 00:05:00 北京时间'
    assert report_beijing_time(datetime(2026, 10, 15, 16, 5)) == '2026-10-16 00:05:00 北京时间'
    data = snapshot()
    data['questions'] = []
    data['citations']['items'] = []
    data['recommendations'] = []
    data['basis']['last_data_at'] = None
    assert '证据最新时间：暂无' in build_project_report_html(data)
