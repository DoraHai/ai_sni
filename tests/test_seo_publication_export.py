from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
from PIL import Image

from app.seo_publication_export import (DEFAULT_PUBLICATION_LIST_TEMPLATE, build_publication_list_workbook,
                                        capture_fields, month_bounds, publication_summary)


def _capture(id, status, source="auto", error_code=None):
    return SimpleNamespace(id=id, status=status, source=source, error_code=error_code,
                           captured_at=datetime(2026, 9, id, tzinfo=timezone.utc), storage_key=f"{id}.png")


def test_month_boundaries_and_capture_rules():
    assert month_bounds("2026-09") == (datetime(2026, 8, 31, 16), datetime(2026, 9, 30, 16))
    assert month_bounds("2026-10")[0] == datetime(2026, 9, 30, 16)
    assert capture_fields([])["capture_status"] == "无截图"
    assert capture_fields([_capture(1, "succeeded")])["capture_status"] == "自动截图"
    assert capture_fields([_capture(1, "succeeded", "manual")])["capture_status"] == "人工上传"
    assert capture_fields([_capture(2, "failed", error_code="timeout")])["capture_status"] == "失败：截图超时"
    selected = capture_fields([_capture(1, "succeeded", "manual"), _capture(2, "failed", error_code="captcha_page")])
    assert selected["capture_status"] == "人工上传"
    assert selected["image_key"] == "1.png"
    assert "人机验证" in selected["capture_note"]


def test_workbook_template_images_bad_image_and_summary():
    image = Image.new("RGB", (300, 500), "#abcdef")
    source = BytesIO()
    image.save(source, "PNG")
    rows = [
        {"platform": "知乎", "title": "=危险标题", "keywords": "", "page_url": "https://example.com/1",
         "published_at": datetime(2026, 9, 30, 15, 59), "capture_status": "自动截图",
         "image_key": "good", "captured_at": datetime(2026, 9, 30, 15, 59), "notes": "", "capture_kind": "auto"},
        {"platform": "官网", "title": "无图", "keywords": "", "page_url": "", "published_at": None,
         "capture_status": "人工上传", "image_key": "bad", "captured_at": None, "notes": "", "capture_kind": "manual"},
    ]
    summary = publication_summary(rows, "测试站", "2026-09", datetime(2026, 10, 1, tzinfo=timezone.utc))
    data = build_publication_list_workbook(rows, summary, image_loader=lambda key: source.getvalue() if key == "good" else b"bad")
    workbook = load_workbook(BytesIO(data))
    detail = workbook["发布明细"]
    assert [cell.value for cell in detail[1]] == [column["title"] for column in DEFAULT_PUBLICATION_LIST_TEMPLATE]
    assert detail["C2"].value == "'=危险标题"
    assert detail["D2"].value is None
    assert detail["F2"].value == "2026-09-30 23:59"
    assert detail["G3"].value == "失败：截图图片不可用"
    assert len(detail._images) == 1
    assert detail._images[0].width == 200
    assert workbook["汇总"]["B4"].value == 2
    assert workbook["汇总"]["B6"].value == 0.5
    assert workbook["汇总"]["B7"].value == 1
    assert workbook["汇总"]["B8"].value == 0
    assert workbook["汇总"]["B9"].value == 1


def test_custom_template_controls_column_order_and_title():
    template = ({"key": "title", "title": "客户标题", "width": 25, "type": "text"},
                {"key": "number", "title": "行号", "width": 8, "type": "text"})
    rows = [{"title": "文章", "platform": "官网", "capture_kind": "none"}]
    summary = publication_summary(rows, "站点", "2026-09", datetime(2026, 10, 1, tzinfo=timezone.utc))
    detail = load_workbook(BytesIO(build_publication_list_workbook(rows, summary, template)))["发布明细"]
    assert [cell.value for cell in detail[1]] == ["客户标题", "行号"]
    assert [cell.value for cell in detail[2]] == ["文章", 1]


def test_empty_publication_coverage_is_not_a_measured_zero():
    summary = publication_summary([], "站点", "2026-09", datetime(2026, 10, 1))
    workbook = load_workbook(BytesIO(build_publication_list_workbook([], summary)))
    assert workbook["汇总"]["B6"].value == "暂无发布记录"
    rows = [{"title": "无截图发布", "capture_kind": "none"}]
    summary = publication_summary(rows, "站点", "2026-09", datetime(2026, 10, 1))
    workbook = load_workbook(BytesIO(build_publication_list_workbook(rows, summary)))
    assert workbook["汇总"]["B6"].value == 0


def test_wrapped_titles_grow_rows_and_preserve_screenshot_height():
    source = BytesIO()
    Image.new("RGB", (200, 240), "white").save(source, "PNG")
    rows = [{"title": "长中文标题" * 30 + "\n第二行", "notes": "备注" * 60, "capture_kind": "none"},
            {"title": "短标题", "image_key": "image", "capture_kind": "auto"},
            {"title": "极长中文标题" * 100, "image_key": "image", "capture_kind": "auto"}]
    summary = publication_summary(rows, "站点", "2026-09", datetime(2026, 10, 1))
    workbook = load_workbook(BytesIO(build_publication_list_workbook(rows, summary, image_loader=lambda _: source.getvalue())))
    detail = workbook["发布明细"]
    assert detail["C2"].value == rows[0]["title"]
    assert detail["C2"].alignment.wrap_text
    assert detail.row_dimensions[2].height >= 150
    assert detail.row_dimensions[3].height >= 185
    assert detail.row_dimensions[4].height == 409
    assert len(detail._images) == 2
    # A narrower customer template needs more lines for the same CJK title.
    template = ({"key": "title", "title": "标题", "width": 20, "type": "text"},)
    narrow = load_workbook(BytesIO(build_publication_list_workbook(rows[:1], summary, template)))
    assert narrow["发布明细"].row_dimensions[2].height > detail.row_dimensions[2].height
